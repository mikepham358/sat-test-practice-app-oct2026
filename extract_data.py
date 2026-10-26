import openpyxl
import json
import re
import os

EXCEL_PATH = "/Users/mike/Library/CloudStorage/GoogleDrive-mikepham358@gmail.com/My Drive/A Vinschool by Mike Pham/26-27 Grand Plan /Study Plan/Module /Template_-_October_2026-v5.xlsx"
PDF_PATH = "/Users/mike/Library/CloudStorage/GoogleDrive-mikepham358@gmail.com/My Drive/A Vinschool by Mike Pham/SAT Coach/Real Tests/Oct 2026/B.pdf"
OUTPUT_DIR = os.path.dirname(os.path.abspath(__file__))
OUTPUT_JSON = os.path.join(OUTPUT_DIR, "database.json")

# Verified PDF page mappings for each question
PDF_PAGE_MAPPING = {
    "RW_Module1": {
        1: 3, 2: 3, 3: 3, 4: 3,
        5: 4, 6: 4,
        7: 5, 8: 5,
        9: 6, 10: 6,
        11: 7,
        12: 8,
        13: 9, 14: 9, 15: 9,
        16: 10, 17: 10, 18: 10,
        19: 11, 20: 11, 21: 11,
        22: 12, 23: 12, 24: 12,
        25: 13, 26: 13,
        27: 14
    },
    "RW_Module2": {
        1: 16, 2: 16, 3: 16, 4: 16,
        5: 17, 6: 17,
        7: 18, 8: 18,
        9: 19, 10: 19,
        11: 20, 12: 20,
        13: 21, 14: 21, 15: 21,
        16: 22, 17: 22, 18: 22,
        19: 23, 20: 23, 21: 23,
        22: 24, 23: 24, 24: 24, 25: 24,
        26: 25, 27: 25
    },
    "Math_Module1": {
        1: 29, 2: 29,
        3: 30, 4: 30, 5: 30, 6: 30, 7: 30,
        8: 31, 9: 31, 10: 31,
        11: 32, 12: 32, 13: 32,
        14: 33, 15: 33, 16: 33, 17: 33,
        18: 34, 19: 34, 20: 34, 21: 34,
        22: 35
    },
    "Math_Module2": {
        1: 38, 2: 38, 3: 38, 4: 38,
        5: 39, 6: 39, 7: 39, 8: 39,
        9: 40, 10: 40, 11: 40,
        12: 41, 13: 41,
        14: 42, 15: 42, 16: 42,
        17: 43, 18: 43, 19: 43,
        20: 44, 21: 44, 22: 44
    }
}

MODULE_INFO = {
    "RW_Module1": {
        "id": "rw1",
        "code": "RW_Module1",
        "title": "Reading & Writing: Module 1",
        "section": "Reading and Writing",
        "module_num": 1,
        "time_minutes": 32,
        "question_count": 27,
        "pdf_start_page": 2
    },
    "RW_Module2": {
        "id": "rw2",
        "code": "RW_Module2",
        "title": "Reading & Writing: Module 2",
        "section": "Reading and Writing",
        "module_num": 2,
        "time_minutes": 32,
        "question_count": 27,
        "pdf_start_page": 15
    },
    "Math_Module1": {
        "id": "math1",
        "code": "Math_Module1",
        "title": "Math: Module 1",
        "section": "Math",
        "module_num": 1,
        "time_minutes": 35,
        "question_count": 22,
        "pdf_start_page": 27
    },
    "Math_Module2": {
        "id": "math2",
        "code": "Math_Module2",
        "title": "Math: Module 2",
        "section": "Math",
        "module_num": 2,
        "time_minutes": 35,
        "question_count": 22,
        "pdf_start_page": 36
    }
}

def parse_options(options_raw):
    if not options_raw:
        return []
    lines = str(options_raw).strip().split('\n')
    parsed = []
    current_key = None
    current_val = []
    
    for line in lines:
        line_s = line.strip()
        match = re.match(r'^([A-D])[\.\:\)]\s*(.*)$', line_s)
        if match:
            if current_key:
                parsed.append({"key": current_key, "text": " ".join(current_val).strip()})
            current_key = match.group(1)
            current_val = [match.group(2)]
        else:
            if current_key:
                current_val.append(line_s)
            elif line_s:
                current_val.append(line_s)
    if current_key:
        parsed.append({"key": current_key, "text": " ".join(current_val).strip()})
    
    if not parsed:
        for opt in ['A', 'B', 'C', 'D']:
            pattern = rf'{opt}[\.\:\)]\s*([^\n]+)'
            m = re.search(pattern, str(options_raw))
            if m:
                parsed.append({"key": opt, "text": m.group(1).strip()})
    return parsed

def extract_hints(formula_hint, formula_tip):
    t1 = ""
    t2 = ""
    t3 = ""
    strategic_tip = ""
    
    if formula_hint:
        f_str = str(formula_hint)
        m3 = re.search(r'Tầng 3[^"]*"\s*&\s*"([^"]+)"', f_str)
        if m3:
            t3 = m3.group(1).strip()
        m2 = re.search(r'Tầng 2[^"]*"\s*&\s*"([^"]+)"', f_str)
        if m2:
            t2 = m2.group(1).strip()
        m1 = re.search(r'Tầng 1[^"]*"\s*&\s*"([^"]+)"', f_str)
        if m1:
            t1 = m1.group(1).strip()
            
    if formula_tip:
        f_tip_str = str(formula_tip)
        parts = re.findall(r'"([^"]*)"', f_tip_str)
        clean_parts = [p for p in parts if not any(x in p for x in ['Đúng', 'Sai', '🔒', 'Trả lời đúng', 'Lần thử', '💡'])]
        if clean_parts:
            strategic_tip = clean_parts[0].strip()
            
    return t1, t2, t3, strategic_tip

def extract_all():
    print(f"Loading workbook from: {EXCEL_PATH}")
    wb = openpyxl.load_workbook(EXCEL_PATH, data_only=False)
    
    # 1. Parse Overview_Strategies
    strategies = []
    if 'Overview_Strategies' in wb.sheetnames:
        ws_strat = wb['Overview_Strategies']
        for r in range(5, ws_strat.max_row + 1):
            category = ws_strat.cell(r, 1).value
            domain = ws_strat.cell(r, 2).value
            core_rules = ws_strat.cell(r, 3).value
            traps = ws_strat.cell(r, 4).value
            if category or domain:
                strategies.append({
                    "category": str(category or "").strip(),
                    "domain": str(domain or "").strip(),
                    "core_rules": str(core_rules or "").strip(),
                    "traps": str(traps or "").strip()
                })
                
    # 2. SAT Error Log Tracker starts clean (0 default logs)
    initial_error_logs = []
                
    # 3. Parse all 4 module sheets
    questions = []
    global_id = 1
    modules_output = []
    
    for mod_code, info in MODULE_INFO.items():
        if mod_code not in wb.sheetnames:
            continue
        ws = wb[mod_code]
        mod_q_ids = []
        
        for r in range(5, ws.max_row + 1):
            q_num_str = ws.cell(r, 1).value
            topic_str = ws.cell(r, 2).value
            passage_str = ws.cell(r, 3).value
            prompt_str = ws.cell(r, 4).value
            options_raw = ws.cell(r, 5).value
            formula_l1 = ws.cell(r, 7).value
            hint_formula = ws.cell(r, 15).value if ws.max_column >= 15 else None
            tip_formula = ws.cell(r, 16).value if ws.max_column >= 16 else None
            
            if not q_num_str:
                continue
                
            m_num = re.search(r'\d+', str(q_num_str))
            q_int = int(m_num.group(0)) if m_num else 1
            
            # Extract correct answer
            ans_match = re.search(r'="([A-D0-9\/\.\-]+)"', str(formula_l1))
            correct_ans = ans_match.group(1) if ans_match else ''
            
            # Extract 3 progressive hint tiers + strategic tip
            t1, t2, t3, strat_tip = extract_hints(hint_formula, tip_formula)
            
            # Fallbacks if somehow empty
            t1_clean = t1 or "Xác định từ khóa trọng tâm và logic ngữ cảnh của câu hỏi."
            t2_clean = t2 or "Áp dụng phương pháp loại trừ và phân tích cấu trúc đối chiếu."
            t3_clean = t3 or "Chốt phương pháp: So sánh trực tiếp với giả thiết cốt lõi trong đề."
            strat_clean = strat_tip or "Nắm vững quy tắc và công thức toán / ngữ pháp College Board."
            
            # Determine if multiple choice or grid-in
            is_grid_in = (correct_ans not in ['A', 'B', 'C', 'D'])
            options = [] if is_grid_in else parse_options(options_raw)
            
            # PDF page
            page = PDF_PAGE_MAPPING.get(mod_code, {}).get(q_int, info["pdf_start_page"])
            
            q_obj = {
                "id": global_id,
                "master_id": f"M{global_id}",
                "q_number": q_int,
                "q_label": str(q_num_str).strip(),
                "module_code": mod_code,
                "module_title": info["title"],
                "section": info["section"],
                "module_num": info["module_num"],
                "topic": str(topic_str or "").strip(),
                "passage": str(passage_str or "").strip(),
                "prompt": str(prompt_str or "").strip(),
                "options": options,
                "options_raw": str(options_raw or "").strip(),
                "is_grid_in": is_grid_in,
                "correct_answer": correct_ans,
                "tier1_hint": t1_clean,
                "tier2_hint": t2_clean,
                "tier3_hint": t3_clean,
                "strategic_tip": strat_clean,
                # Backward compatibility aliases
                "level1_tip": t1_clean,
                "level2_hint": t2_clean,
                "level3_solution": t3_clean,
                "pdf_page": page
            }
            questions.append(q_obj)
            mod_q_ids.append(global_id)
            global_id += 1
            
        mod_copy = dict(info)
        mod_copy["question_ids"] = mod_q_ids
        mod_copy["question_count"] = len(mod_q_ids)
        modules_output.append(mod_copy)
        
    output_data = {
        "test_title": "Digital SAT Practice Test — October 2026 (INT-B)",
        "pdf_path": PDF_PATH,
        "pdf_drive_preview": "https://drive.google.com/file/d/1Qiu34_dltehpEVmfHvOivjAvf8Zkd77b/preview",
        "pdf_drive_view": "https://drive.google.com/file/d/1Qiu34_dltehpEVmfHvOivjAvf8Zkd77b/view?usp=drive_link",
        "excel_path": EXCEL_PATH,
        "total_questions": len(questions),
        "modules": modules_output,
        "questions": questions,
        "initial_error_logs": initial_error_logs,
        "strategies": strategies
    }
    
    with open(OUTPUT_JSON, "w", encoding="utf-8") as f:
        json.dump(output_data, f, ensure_ascii=False, indent=2)
        
    print(f"Extraction complete! Saved to {OUTPUT_JSON}")
    print(f"Total questions extracted: {len(questions)}")
    for m in modules_output:
        print(f"  {m['title']}: {m['question_count']} questions (pages {min(PDF_PAGE_MAPPING[m['code']].values())}-{max(PDF_PAGE_MAPPING[m['code']].values())})")

if __name__ == "__main__":
    extract_all()
