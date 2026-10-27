import os
import re
import json
import fitz
import openpyxl

APP_DIR = "/Users/mike/Library/CloudStorage/GoogleDrive-mikepham358@gmail.com/My Drive/A Vinschool by Mike Pham/SAT Coach/BluebookOct2026App"
PDF_PATH = "/Users/mike/Library/CloudStorage/GoogleDrive-mikepham358@gmail.com/My Drive/A Vinschool by Mike Pham/SAT Coach/Real Tests/Oct 2026/B.pdf"
DB_PATH = os.path.join(APP_DIR, "database.json")
STATIC_DB_PATH = os.path.join(APP_DIR, "static", "database.json")
TEMPLATE_PATH = os.path.join(APP_DIR, "template.xlsx")
STATIC_TEMPLATE_PATH = os.path.join(APP_DIR, "static", "template.xlsx")
MASTER_TEMPLATE_PATH = "/Users/mike/Library/CloudStorage/GoogleDrive-mikepham358@gmail.com/My Drive/A Vinschool by Mike Pham/26-27 Grand Plan /Study Plan/Module /Template - October 2026.xlsx"

def extract_rw_from_pdf():
    doc = fitz.open(PDF_PATH)
    
    def get_module_text(pstart, pend):
        mod_text = ""
        for pno in range(pstart, pend):
            text = doc[pno].get_text()
            cleaned_lines = []
            for line in text.splitlines():
                l = line.strip()
                if re.search(r"Digital SAT Practice Test", l):
                    continue
                if re.search(r"sattest\.hocsat\.com\s*\|\s*Page", l):
                    continue
                if re.search(r"Section \d+, Module \d+:", l):
                    continue
                cleaned_lines.append(line)
            mod_text += "\n" + "\n".join(cleaned_lines)
        return mod_text

    def format_passage(raw_passage):
        lines = [l.strip() for l in raw_passage.splitlines() if l.strip()]
        formatted_paras = []
        curr_para = []
        
        for line in lines:
            if line.startswith("•") or line.startswith("-") or line.startswith("Text 1") or line.startswith("Text 2") or line.startswith("Table") or line.startswith("©"):
                if curr_para:
                    formatted_paras.append(" ".join(curr_para))
                    curr_para = []
                formatted_paras.append(line)
            else:
                curr_para.append(line)
        if curr_para:
            formatted_paras.append(" ".join(curr_para))
            
        return "\n\n".join(formatted_paras)

    def parse_q_block(qtext):
        opt_pattern = r"(?:^|\n)\s*([A-D])\.\s*(.+?)(?=(?:\n\s*[A-D]\.\s*)|$)"
        opts_matches = list(re.finditer(opt_pattern, qtext, re.DOTALL))
        if len(opts_matches) != 4:
            raise ValueError(f"Expected 4 options, found {len(opts_matches)}")
            
        first_opt_start = opts_matches[0].start()
        pre_opts = qtext[:first_opt_start].strip()
        
        opts = []
        for m in opts_matches:
            key = m.group(1)
            text = " ".join(m.group(2).split()).strip()
            opts.append({"key": key, "text": text})
            
        m1 = re.search(r"(The student wants to.+?\?)$", pre_opts, re.DOTALL)
        if m1:
            prompt = m1.group(1).strip()
            passage = pre_opts[:m1.start()].strip()
        else:
            m2 = re.search(r"((?:Which (?:choice|finding|quotation|statement)|According to|Information in).+?\?)$", pre_opts, re.DOTALL)
            if m2:
                prompt = m2.group(1).strip()
                passage = pre_opts[:m2.start()].strip()
            else:
                prompt = ""
                passage = pre_opts
                
        prompt_clean = " ".join(prompt.split()).strip()
        passage_clean = format_passage(passage)
        return passage_clean, prompt_clean, opts

    modules_data = {}
    for mod_name, pstart, pend in [("RW_Module1", 2, 14), ("RW_Module2", 15, 26)]:
        text = get_module_text(pstart, pend)
        q_blocks = re.split(r"\n\s*(\d{1,2})\s*\n", text)
        qs = {}
        for i in range(1, len(q_blocks), 2):
            qnum = int(q_blocks[i])
            passage, prompt, opts = parse_q_block(q_blocks[i+1])
            qs[qnum] = {
                "passage": passage,
                "prompt": prompt,
                "opts": opts,
                "options_raw": "\n".join([f"{o['key']}. {o['text']}" for o in opts])
            }
        modules_data[mod_name] = qs
    return modules_data

def update_database(rw_data):
    print("Updating database.json...")
    with open(DB_PATH, "r", encoding="utf-8") as f:
        db = json.load(f)

    updated_count = 0
    for q in db["questions"]:
        mod = q["module_code"]
        qnum = q["q_number"]
        if mod in rw_data and qnum in rw_data[mod]:
            pdata = rw_data[mod][qnum]
            q["options"] = pdata["opts"]
            q["options_raw"] = pdata["options_raw"]
            if pdata["passage"]:
                q["passage"] = pdata["passage"]
            if pdata["prompt"]:
                q["prompt"] = pdata["prompt"]
            updated_count += 1

    with open(DB_PATH, "w", encoding="utf-8") as f:
        json.dump(db, f, ensure_ascii=False, indent=2)

    with open(STATIC_DB_PATH, "w", encoding="utf-8") as f:
        json.dump(db, f, ensure_ascii=False, indent=2)

    print(f"Updated {updated_count} questions in database.json and static/database.json.")

def update_excel_files(rw_data):
    for path in [TEMPLATE_PATH, STATIC_TEMPLATE_PATH, MASTER_TEMPLATE_PATH]:
        if not os.path.exists(path):
            print(f"Warning: {path} not found, skipping.")
            continue
        print(f"Updating Excel: {path}...")
        wb = openpyxl.load_workbook(path)
        
        # 1. Update RW_Module1 and RW_Module2
        for mod in ["RW_Module1", "RW_Module2"]:
            if mod in wb.sheetnames:
                ws = wb[mod]
                for qnum, data in rw_data[mod].items():
                    r = 4 + qnum
                    if data["passage"]:
                        ws.cell(r, 3).value = data["passage"]
                    if data["prompt"]:
                        ws.cell(r, 4).value = data["prompt"]
                    ws.cell(r, 5).value = data["options_raw"]
                    
        # 2. Update Master_All_98_Questions
        if "Master_All_98_Questions" in wb.sheetnames:
            ws_master = wb["Master_All_98_Questions"]
            # RW1: rows 5 to 31
            for qnum, data in rw_data["RW_Module1"].items():
                r = 4 + qnum
                if data["passage"]:
                    ws_master.cell(r, 4).value = data["passage"]
                if data["prompt"]:
                    ws_master.cell(r, 5).value = data["prompt"]
                ws_master.cell(r, 6).value = data["options_raw"]
            # RW2: rows 32 to 58
            for qnum, data in rw_data["RW_Module2"].items():
                r = 4 + 27 + qnum
                if data["passage"]:
                    ws_master.cell(r, 4).value = data["passage"]
                if data["prompt"]:
                    ws_master.cell(r, 5).value = data["prompt"]
                ws_master.cell(r, 6).value = data["options_raw"]
                
        wb.save(path)
        print(f"Saved {path} successfully.")

def verify_all():
    print("\n--- Running Verification ---")
    with open(DB_PATH, "r", encoding="utf-8") as f:
        db = json.load(f)

    db_ellipses = []
    for q in db["questions"]:
        for opt in q.get("options", []):
            txt = opt.get("text", "")
            if "..." in txt or "…" in txt:
                db_ellipses.append((q["module_code"], q["q_number"], opt["key"], txt))

    print(f"database.json options with ellipsis: {len(db_ellipses)}")
    for item in db_ellipses:
        print(f"  {item[0]} Q{item[1]} ({item[2]}): {item[3]}")

    for path in [TEMPLATE_PATH, STATIC_TEMPLATE_PATH, MASTER_TEMPLATE_PATH]:
        wb = openpyxl.load_workbook(path, data_only=True)
        excel_ellipses = []
        for mod in ["RW_Module1", "RW_Module2"]:
            ws = wb[mod]
            for r in range(5, ws.max_row + 1):
                raw = str(ws.cell(r, 5).value or "")
                if "..." in raw or "…" in raw:
                    excel_ellipses.append((mod, ws.cell(r, 1).value, raw))
        if "Master_All_98_Questions" in wb.sheetnames:
            ws = wb["Master_All_98_Questions"]
            for r in range(5, 59): # RW1 and RW2
                raw = str(ws.cell(r, 6).value or "")
                if "..." in raw or "…" in raw:
                    excel_ellipses.append(("Master", ws.cell(r, 1).value, raw))
        print(f"{os.path.basename(path)} options with ellipsis: {len(excel_ellipses)}")

if __name__ == "__main__":
    rw_data = extract_rw_from_pdf()
    print("Extracted RW data from PDF successfully.")
    update_database(rw_data)
    update_excel_files(rw_data)
    verify_all()
