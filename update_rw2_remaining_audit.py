import json
import openpyxl

AUDIT_FIXES = {
    6: {
        "id": 33,
        "master_id": "M33",
        "q_number": 6,
        "module_code": "RW_Module2",
        "topic": "Text Structure and Purpose (Overall Text Structure)",
        "difficulty": "Medium",
        "passage": "With examples ranging from Dean’s Blue Hole in the Bahamas to Watling’s Blue Hole near San Salvador Island, blue holes are large underwater sinkholes: pits filled with dark blue water that contrasts with lighter blue shallower waters surrounding them. The formation of blue holes typically involves several natural processes, including the dissolution of limestone or other carbonate rocks. Currently, the deepest blue hole on record is the Taam Ja’ Blue Hole in Chetumal Bay, Mexico, which is reported by Laura Carrillo and colleagues as extending at least 420 meters below sea level.",
        "prompt": "Which choice best describes the overall structure of the text?",
        "options_raw": "A. It describes the appearance of a particular type of sinkhole, compares the formations of two sinkholes of that type, and then names the researchers who discovered the deepest known sinkhole of that type.\nB. It compares a deep sinkhole to a shallower one, details the processes by which both sinkholes formed, and emphasizes the remarkable depth of the deeper sinkhole.\nC. It explains how a particular sinkhole formed, evaluates that sinkhole’s unusual appearance, and names a research team that is studying the sinkhole.\nD. It defines a kind of sinkhole, describes its typical appearance and formation, and provides information about the deepest known sinkhole.",
        "correct_answer": "D",
        "princeton_strategy": "Structure & Purpose - Text Mapping",
        "tier1_hint": "Lập sơ đồ cấu trúc đoạn văn qua 3 phần: Câu 1 đưa ra định nghĩa và diện mạo của hố xanh (blue holes); câu 2 mô tả quá trình hình thành địa chất (formation); câu 3 cung cấp thông tin về hố xanh sâu nhất hiện nay (Taam Ja’ Blue Hole).",
        "tier2_hint": "Đối chiếu sơ đồ này với 4 phương án: Tìm phương án bao quát được cả 3 bước: định nghĩa (defines), miêu tả diện mạo & hình thành (appearance and formation), và thông tin về hố sâu nhất (deepest known sinkhole).",
        "tier3_hint": "Phương án D khớp chính xác 100% với trình tự phát triển của văn bản. Chọn D.",
        "strategic_tip": "Mẹo Princeton Review: Đừng để bị lừa bởi các chi tiết riêng lẻ! Luôn tóm tắt vai trò của từng câu trước khi nhìn vào các phương án.",
        "level1_tip": "Chia đoạn văn làm 3 phần: định nghĩa -> quá trình hình thành -> hố sâu nhất kỷ lục.",
        "level2_hint": "Loại bỏ phương án có sự so sánh giữa 2 hố cụ thể (A, B) hoặc chỉ tập trung vào 1 hố (C).",
        "level3_solution": "Chọn D. It defines a kind of sinkhole, describes its typical appearance and formation, and provides information about the deepest known sinkhole.",
        "target_clue": "Câu 1: Định nghĩa hố xanh ('blue holes are large underwater sinkholes...'). Câu 2: Quá trình hình thành ('formation... involves natural processes'). Câu 3: Kỷ lục hố sâu nhất ('deepest blue hole on record is the Taam Ja’...').",
        "solution_steps": "1. Phân tích cấu trúc đoạn văn:\n   - Đoạn 1: Định nghĩa hố xanh là gì và miêu tả đặc điểm màu sắc đặc trưng của nó.\n   - Đoạn 2: Trình bày cơ chế hình thành tự nhiên (hòa tan đá vôi/carbonate).\n   - Đoạn 3: Cung cấp số liệu về hố xanh Taam Ja' sâu nhất thế giới (420m).\n2. Đối chiếu phương án:\n   - A sai vì đoạn văn không so sánh sự hình thành của hai hố cụ thể.\n   - B sai vì văn bản không so sánh một hố sâu với một hố nông.\n   - C sai vì bài không chỉ nói về một hố duy nhất từ đầu đến cuối.\n   - D đúng hoàn toàn: định nghĩa hố sụt (defines a kind of sinkhole), miêu tả diện mạo & hình thành (appearance and formation), và nêu hố sâu nhất (deepest known sinkhole).",
        "pooe_breakdown": "• Đáp án ĐÚNG: D. It defines a kind of sinkhole, describes its typical appearance and formation, and provides information about the deepest known sinkhole. — Phản ánh chính xác và toàn diện cả 3 bước cấu trúc của đoạn trích.\n• Loại A. It describes... compares the formations of two sinkholes — Sai chi tiết; bài không so sánh cơ chế hình thành của hai hố.\n• Loại B. It compares a deep sinkhole to a shallower one — Sai bố cục; bài không lập bảng so sánh hố sâu và hố cạn.\n• Loại C. It explains how a particular sinkhole formed — Sai phạm vi; 2 câu đầu nói về hố xanh nói chung chứ không phải một hố cá biệt.",
        "pro_tip": "Mẹo Princeton Review: Đối với câu hỏi cấu trúc văn bản (Text Structure), hãy viết ra nháp chuỗi động từ chính (Define -> Explain formation -> Detail record) để đối chiếu nhanh.",
        "option_explanations": {
            "A": "Loại A — Sai chi tiết cấu trúc. Tác giả chỉ nêu tên hố Dean và Watling làm ví dụ cho sự phân bố địa lý, hoàn toàn không so sánh cơ chế hình thành giữa hai hố này.",
            "B": "Loại B — Sai bố cục tổng thể. Văn bản không xây dựng theo cấu trúc so sánh - đối chiếu giữa một hố sâu và một hố nông.",
            "C": "Loại C — Sai phạm vi. Bài viết mở đầu bằng việc định nghĩa hiện tượng hố xanh trên diện rộng chứ không chỉ giải thích sự hình thành của một hố cá biệt.",
            "D": "ĐÁP ÁN ĐÚNG: D — Khớp 100% với tiến trình văn bản: Định nghĩa hố xanh (defines a kind of sinkhole), mô tả diện mạo và quá trình tạo thành (appearance and formation), và kết lại bằng kỷ lục hố sâu nhất thế giới (deepest known sinkhole)."
        }
    },
    8: {
        "id": 35,
        "master_id": "M35",
        "q_number": 8,
        "module_code": "RW_Module2",
        "topic": "Central Ideas and Details (Main Idea)",
        "difficulty": "Hard",
        "passage": "Primarily encountered by modern readers as inert written texts, the vernacular riddles of early medieval England contain “implicit vestiges” of an interactive performance context, scholar Jill Frederick argues. Locative phrases (e.g., “here” and “now”) scattered throughout the riddles indicate the space shared by their speakers and audiences; the common refrain “say what I am called” invites a verbalized response; and the conflation of many of the riddles’ speakers with their solutions (typically animals, objects, or natural phenomena) suggests a form of dramatic impersonation perhaps meant to entertain during public gatherings.",
        "prompt": "Which choice best states the main idea of the text?",
        "options_raw": "A. It is reasonable to conclude based on textual evidence from the vernacular riddles of early medieval England that they were mainly performed at large public gatherings, though documentation of the locations in which they were performed is unlikely to exist.\nB. Linguistic evidence found within the vernacular riddles of early medieval England suggests that they may have been intended to be performed in communal settings, though this isn’t how they are generally experienced by readers today.\nC. People of early medieval England likely found that era’s vernacular riddles to be more amusing than modern readers do, given that there is substantial evidence to indicate that they are best experienced in the context of an interactive performance rather than exclusively as written texts.\nD. Scholars once assumed that the vernacular riddles of early medieval England were originally intended to be read silently, but newly uncovered historical documents corroborate Jill Frederick’s assertion that they may in fact have been intended primarily for oral performance.",
        "correct_answer": "B",
        "princeton_strategy": "Central Idea - Core Thesis Identification",
        "tier1_hint": "Tìm câu chứa luận điểm cốt lõi (Thesis statement): Ngay câu mở đầu, học giả Jill Frederick khẳng định dù ngày nay độc giả tiếp cận câu đố thời trung cổ dưới dạng 'văn bản viết tĩnh' (inert written texts), chúng vẫn chứa đựng các 'dấu tích ngầm' của bối cảnh biểu diễn tương tác (interactive performance).",
        "tier2_hint": "Các câu sau đưa ra 3 bằng chứng ngôn ngữ: từ chỉ không gian (here, now), lời mời đáp lại (say what I am called), và sự nhập vai kịch nghệ. Điều này chứng minh chúng được tạo ra để biểu diễn trong cộng đồng.",
        "tier3_hint": "Phương án B tóm lược hoàn hảo: Các bằng chứng ngôn ngữ cho thấy câu đố được tạo ra để biểu diễn trong không gian cộng đồng, dù ngày nay chúng không còn được trải nghiệm theo cách đó. Chọn B.",
        "strategic_tip": "Mẹo Princeton Review: Đối với Main Idea, hãy cảnh giác với các phương án bịa ra giả định ngoài bài như 'newly uncovered historical documents' (D) hoặc so sánh độ hài hước (C).",
        "level1_tip": "Luận điểm chính: Bằng chứng ngôn ngữ trong câu đố trung cổ cho thấy chúng vốn được dùng để diễn xướng trong cộng đồng.",
        "level2_hint": "Loại phương án suy diễn tài liệu mới (D), so sánh cảm xúc (C) hoặc khẳng định quá đà về địa điểm (A).",
        "level3_solution": "Chọn B. Linguistic evidence found within the vernacular riddles...",
        "target_clue": "Luận điểm mở đầu: 'contain implicit vestiges of an interactive performance context' kết hợp đối chiếu với trải nghiệm hiện đại 'primarily encountered by modern readers as inert written texts'.",
        "solution_steps": "1. Câu chủ đề: Các câu đố dân gian thời sơ kỳ trung cổ Anh dù nay được đọc như văn bản chết, nhưng chứa các dấu hiệu của việc biểu diễn tương tác.\n2. Các dẫn chứng: Từ ngữ chỉ vị trí ('here', 'now'), câu nhắc đối thoại ('say what I am called'), và sự hóa thân vào đồ vật/con vật.\n3. Khái quát hóa ý chính: Bằng chứng ngôn ngữ chỉ ra các câu đố vốn nhằm mục đích diễn xướng cộng đồng, trái ngược với cách đọc thụ động ngày nay.\n4. Loại A vì khẳng định quá mức ('mainly performed at large public gatherings') và thêm chi tiết tài liệu không tồn tại. Loại C vì so sánh cảm xúc 'more amusing'. Loại D vì bịa ra 'newly uncovered historical documents'. Chọn B.",
        "pooe_breakdown": "• Đáp án ĐÚNG: B. Linguistic evidence found within the vernacular riddles... — Thâu tóm trung thực và trọn vẹn luận điểm của Jill Frederick cùng các minh chứng ngôn ngữ trong bài.\n• Loại A. It is reasonable to conclude... documentation is unlikely to exist — Suy diễn thái quá về quy mô tụ họp và phỏng đoán tài liệu.\n• Loại C. People of early medieval England found riddles more amusing... — So sánh sai lệch về mức độ giải trí giữa hai thời kỳ.\n• Loại D. Scholars once assumed... newly uncovered historical documents — Bẫy sai dữ kiện; bài phân tích câu từ có sẵn chứ không nhắc đến tài liệu lịch sử mới tìm thấy.",
        "pro_tip": "Mẹo Princeton Review: Main Idea phải bao trùm toàn văn bản, không được chứa các chi tiết ngoại suy nằm ngoài phạm vi đoạn trích.",
        "option_explanations": {
            "A": "Loại A — Suy diễn cực đoan. Bài viết không khẳng định câu đố 'chủ yếu diễn ra ở các cuộc tụ họp lớn' và cũng không bàn về việc tài liệu địa điểm có tồn tại hay không.",
            "B": "ĐÁP ÁN ĐÚNG: B — Diễn giải hoàn hảo ý chính của văn bản: Bằng chứng ngôn ngữ trong các câu đố cho thấy chúng vốn được biểu diễn trong không gian cộng đồng, dù ngày nay độc giả chỉ tiếp cận dưới dạng văn bản tĩnh.",
            "C": "Loại C — So sánh không có căn cứ. Văn bản không hề khẳng định người xưa thấy câu đố 'hài hước / thú vị hơn' độc giả thời hiện đại.",
            "D": "Loại D — Bẫy dữ kiện sai thực tế (False Premise). Đoạn văn không hề nhắc đến việc 'các tài liệu lịch sử mới được phát hiện' chứng minh cho khẳng định của Jill Frederick."
        }
    },
    9: {
        "id": 36,
        "master_id": "M36",
        "q_number": 9,
        "module_code": "RW_Module2",
        "topic": "Inferences (Supporting Inferences / Logic Deduction)",
        "difficulty": "Hard",
        "passage": "In the Mexican dialect of Spanish, the consonant s is pronounced with greater tension in the tongue than it is in Castilian, the dominant dialect in Spain, resulting in a sharper sound. Margarita Hidalgo, a linguist at San Diego State University, argues that Mexican Spanish acquired this pronunciation from Nahuatl, the language of the Nahua people in central Mexico: as the Spanish colonizers of Mexico incorporated Nahuatl words into their speech, and as more and more Nahuas spoke Spanish as a second language, the Nahuatl pronunciation of the s consonant replaced the Castilian pronunciation.",
        "prompt": "Information in the text best supports which statement about the consonant s?",
        "options_raw": "A. The pronunciation of this consonant is the principal difference between the Mexican and Castilian dialects of Spanish.\nB. This consonant appears in a higher percentage of Nahuatl vocabulary than of Castilian vocabulary.\nC. The pronunciation of this consonant in Nahuatl may have been affected by its pronunciation in Mexican Spanish.\nD. The pronunciation of this consonant requires greater tension in the tongue in Nahuatl than in Castilian.",
        "correct_answer": "D",
        "princeton_strategy": "Inferences - Strict Textual Deduction",
        "tier1_hint": "Theo dõi chuỗi logic trong bài: 1) Trong tiếng TBN Mexico, âm 's' được phát âm với độ căng của lưỡi lớn hơn so với tiếng Castilian (Tây Ban Nha chính quốc). 2) Tiếng TBN Mexico có được cách phát âm này là do tiếp nhận từ tiếng Nahuatl (cách phát âm của tiếng Nahuatl thay thế cách phát âm của Castilian).",
        "tier2_hint": "Từ hai dữ kiện trên, ta có thể suy ra điều gì về độ căng của lưỡi khi phát âm 's' trong tiếng Nahuatl so với tiếng Castilian?",
        "tier3_hint": "Bởi vì tiếng TBN Mexico học cách phát âm 's' căng lưỡi từ tiếng Nahuatl, nên trong tiếng Nahuatl, âm 's' chắc chắn đòi hỏi độ căng của lưỡi nhiều hơn so với tiếng Castilian. Chọn D.",
        "strategic_tip": "Mẹo Princeton Review: Đối với câu hỏi Inference, đáp án đúng phải là suy luận tất yếu trực tiếp từ văn bản (A = B và B = C thì A = C), tránh các từ tuyệt đối như 'principal difference' (A).",
        "level1_tip": "Dữ kiện 1: TBN Mexico phát âm 's' căng lưỡi hơn Castilian. Dữ kiện 2: TBN Mexico lấy cách phát âm này từ tiếng Nahuatl.",
        "level2_hint": "Suy luận tất yếu: Tiếng Nahuatl phát âm 's' căng lưỡi hơn tiếng Castilian.",
        "level3_solution": "Chọn D. The pronunciation of this consonant requires greater tension in the tongue in Nahuatl than in Castilian.",
        "target_clue": "Liên kết bắc cầu: 'Mexican Spanish: greater tension than Castilian' + 'Mexican Spanish acquired this pronunciation from Nahuatl' => 'Nahuatl requires greater tension than Castilian'.",
        "solution_steps": "1. Văn bản nêu: Tiếng TBN Mexico phát âm 's' có độ căng lưỡi lớn hơn tiếng Castilian (Tây Ban Nha chính quốc).\n2. Văn bản khẳng định: Tiếng TBN Mexico có được cách phát âm này từ tiếng Nahuatl (phát âm Nahuatl thay thế phát âm Castilian).\n3. Suy luận logic bắc cầu: Vì phát âm của Nahuatl truyền độ căng lưỡi này sang TBN Mexico, nên cách phát âm âm 's' trong tiếng Nahuatl đòi hỏi độ căng lưỡi nhiều hơn so với tiếng Castilian.\n4. Đối chiếu các phương án:\n   - A sai vì 'principal difference' (sự khác biệt cốt lõi nhất) là suy diễn quá đà.\n   - B sai vì bài không nói về tỷ lệ phần trăm từ vựng chứa âm 's'.\n   - C sai vì đảo ngược chiều ảnh hưởng (Nahuatl ảnh hưởng TBN, không phải ngược lại).\n   - D đúng hoàn toàn.",
        "pooe_breakdown": "• Đáp án ĐÚNG: D. The pronunciation of this consonant requires greater tension in the tongue in Nahuatl than in Castilian. — Suy luận bắc cầu chuẩn xác 100% từ dữ kiện văn bản.\n• Loại A. The pronunciation... is the principal difference — Lỗi phóng đại; bài không khẳng định đây là sự khác biệt quan trọng nhất.\n• Loại B. appears in a higher percentage of Nahuatl vocabulary — Lỗi thông tin không có trong bài; không có số liệu tỷ lệ từ vựng.\n• Loại C. in Nahuatl may have been affected by its pronunciation in Mexican Spanish — Lỗi đảo ngược quan hệ nhân quả.",
        "pro_tip": "Mẹo Princeton Review: Inference trong SAT luôn là một bước suy luận tối thiểu, an toàn và có căn cứ trực tiếp từ chữ trong đề, tuyệt đối không đưa ra giả định ngoài luồng.",
        "option_explanations": {
            "A": "Loại A — Lỗi cường điệu hóa. Đoạn văn chỉ nêu đây là một điểm khác biệt về ngữ âm, hoàn toàn không khẳng định đây là 'khác biệt chính yếu nhất' giữa hai phương ngữ.",
            "B": "Loại B — Lỗi thông tin không tồn tại. Không có dữ liệu nào so sánh tần suất xuất hiện của phụ âm 's' trong kho từ vựng của hai ngôn ngữ.",
            "C": "Loại C — Đảo ngược chiều tác động. Tiếng Nahuatl tác động làm thay đổi tiếng Tây Ban Nha Mexico, chứ không phải tiếng Nahuatl bị biến đổi bởi tiếng Tây Ban Nha.",
            "D": "ĐÁP ÁN ĐÚNG: D — Suy luận logic tuyệt đối. Vì tiếng Tây Ban Nha Mexico phát âm âm 's' căng lưỡi hơn tiếng Castilian, và đặc điểm này được du nhập từ tiếng Nahuatl, suy ra việc phát âm âm 's' trong tiếng Nahuatl đòi hỏi độ căng lưỡi lớn hơn trong tiếng Castilian."
        }
    },
    10: {
        "id": 37,
        "master_id": "M37",
        "q_number": 10,
        "module_code": "RW_Module2",
        "topic": "Command of Evidence (Support a Conclusion)",
        "difficulty": "Medium",
        "passage": "Cover crops, such as teff, are planted by many farmers between main growing seasons. While these crops are known for increasing biodiversity, Axel Don and his coauthor argued in a 2015 study that cover crops’ role in supporting soil carbon sequestration deserves more attention. Cover crops remove carbon dioxide (CO2) from the air and convert it to organic compounds; some of these compounds are stored in the soil as soil organic carbon (SOC), which contributes to the buildup of SOC stocks over time. Based on data from numerous farm fields, Don and his coauthor concluded that cover crops contribute meaningfully to soil carbon sequestration and, if widely adopted, could help offset agricultural emissions of atmospheric CO2 and other greenhouse gases.",
        "prompt": "Which finding, if true, would be most consistent with Don and his coauthor’s conclusion?",
        "options_raw": "A. Teff captures considerably more CO2 in fields where it is the sole cover crop compared with fields where it is grown alongside other species of cover crops.\nB. Cover crops generally exhibit reduced carbon sequestration rates in fields exposed to high levels of greenhouse gas emissions compared with fields exposed to low levels of those emissions.\nC. SOC stocks vary more in fields where cover crops have been grown for decades compared with fields where cover crops have only recently been introduced.\nD. SOC stocks tend to be significantly lower in fields without cover crops compared with fields where cover crops have been grown for several years.",
        "correct_answer": "D",
        "princeton_strategy": "Command of Evidence - Hypothesis Testing",
        "tier1_hint": "Xác định kết luận cần ủng hộ: 'cover crops contribute meaningfully to soil carbon sequestration... buildup of SOC stocks over time' (cây che phủ đóng góp ý nghĩa vào việc cô lập carbon trong đất và làm tăng trữ lượng SOC theo thời gian).",
        "tier2_hint": "Bằng chứng củng cố mạnh nhất là một phép so sánh đối chứng (A/B testing): So sánh lượng SOC giữa cánh đồng CÓ trồng cover crops và cánh đồng KHÔNG trồng cover crops.",
        "tier3_hint": "Phương án D chỉ ra rằng trữ lượng SOC ở những cánh đồng không trồng cover crops thấp hơn đáng kể so với những cánh đồng có trồng cover crops nhiều năm. Đây là minh chứng đối chứng hoàn hảo chứng minh hiệu quả của cover crops. Chọn D.",
        "strategic_tip": "Mẹo Princeton Review: Để chứng minh X có tác dụng làm tăng Y, bằng chứng thực nghiệm mạnh nhất luôn là: Nơi KHÔNG CÓ X thì Y thấp hơn đáng kể so với nơi CÓ X.",
        "level1_tip": "Kết luận: Trồng cây che phủ đất làm tăng trữ lượng carbon hữu cơ trong đất (SOC).",
        "level2_hint": "Tìm kết quả đối chứng: Ruộng không trồng cây che phủ có SOC thấp hơn nhiều so với ruộng có trồng.",
        "level3_solution": "Chọn D. SOC stocks tend to be significantly lower in fields without cover crops...",
        "target_clue": "Kết luận nghiên cứu: 'cover crops contribute meaningfully to soil carbon sequestration and... buildup of SOC stocks'. Bằng chứng cần chứng minh sự gia tăng SOC nhờ có cover crops.",
        "solution_steps": "1. Kết luận của Axel Don: Trồng cây che phủ (cover crops) giúp hấp thụ CO2 và tích tụ carbon hữu cơ trong đất (SOC stocks) theo thời gian.\n2. Cần tìm phát hiện thực nghiệm chứng minh nhận định này là đúng.\n3. Phân tích phương án D: Trữ lượng SOC ở các cánh đồng không có cover crops thấp hơn đáng kể so với cánh đồng trồng cover crops nhiều năm -> Chứng minh trực tiếp rằng việc trồng cover crops mang lại sự tích tụ SOC rõ rệt.\n4. Các phương án khác:\n   - A so sánh giữa trồng riêng teff và trồng xen canh, không chứng minh được tác dụng tổng thể của cover crops.\n   - B nói về sự suy giảm khả năng hấp thụ khi nồng độ khí thải cao, làm yếu đi luận điểm.\n   - C nói về độ biến thiên (vary more), không chứng minh được trữ lượng tăng hay giảm.",
        "pooe_breakdown": "• Đáp án ĐÚNG: D. SOC stocks tend to be significantly lower in fields without cover crops... — Cung cấp bằng chứng thực nghiệm đối chứng trực tiếp chứng minh cover crops giúp gia tăng trữ lượng SOC.\n• Loại A. Teff captures considerably more CO2... sole cover crop — So sánh nội bộ giữa các phương pháp trồng, không củng cố kết luận chung về việc tăng carbon.\n• Loại B. exhibit reduced carbon sequestration rates... — Đưa ra thông tin bất lợi về sự sụt giảm hiệu quả khi gặp khí thải.\n• Loại C. SOC stocks vary more... — Nói về biên độ dao động trữ lượng, không chứng minh được lượng tích lũy thực tế.",
        "pro_tip": "Mẹo Princeton Review: Để chứng minh X có tác dụng làm tăng Y, bằng chứng thực nghiệm mạnh nhất luôn là: Nơi KHÔNG CÓ X thì Y thấp hơn đáng kể so với nơi CÓ X.",
        "option_explanations": {
            "A": "Loại A — Lệch trọng tâm. Việc so sánh hiệu quả giữa việc trồng riêng teff hay trồng xen kẽ chỉ là kỹ thuật canh tác, không chứng minh được vai trò tổng thể của cover crops đối với đất.",
            "B": "Loại B — Phản tác dụng. Thông tin tỷ lệ cô lập carbon giảm ở nơi có nhiều khí thải không hỗ trợ cho kết luận về tính hiệu quả của cây che phủ.",
            "C": "Loại C — Không rõ ràng. Mức độ dao động (vary more) của trữ lượng SOC không đồng nghĩa với việc lượng carbon được giữ lại nhiều hơn.",
            "D": "ĐÁP ÁN ĐÚNG: D — Bằng chứng đối chứng mẫu mực. Việc cánh đồng không trồng cover crops có lượng SOC thấp hơn hẳn cánh đồng có trồng cover crops chứng minh trực tiếp rằng cover crops thực sự thúc đẩy quá trình tích tụ carbon hữu cơ trong đất."
        }
    },
    11: {
        "id": 38,
        "master_id": "M38",
        "q_number": 11,
        "module_code": "RW_Module2",
        "topic": "Inferences (Logical Completion)",
        "difficulty": "Hard",
        "passage": "While the brain regions most strongly associated with language tasks have some modularity in that different language tasks activate these areas to varying degrees, comprehension also invokes other areas of the brain more strongly associated with nonlinguistic functions. For example, understanding a description of a person performing physical actions might call on parts of the listener’s brain associated with sensory-motor functions, whereas a highly emotive account will likely activate areas associated with affective processing. Thus, beyond the areas most strongly associated with linguistic functions, ______",
        "prompt": "Which choice most logically completes the text?",
        "options_raw": "A. the specific combination of brain regions involved in comprehending speech is likely contingent on the total context in which the speech is embedded.\nB. the brain regions associated with sensory-motor functions are often involved in general comprehension, even in contexts that don’t involve descriptions of physical movement.\nC. the brain regions most strongly associated with affective processing activate whenever a listener believes an utterance has emotional connotations, even if the speaker did not intend those connotations.\nD. the areas of the brain that activate when processing speech will typically include areas primarily responsible for nonlinguistic functions, although the exact combination of areas may be impossible to determine.",
        "correct_answer": "A",
        "princeton_strategy": "Inferences - Synthesizing Examples into a General Rule",
        "tier1_hint": "Chú ý từ nối kết luận 'Thus, beyond the areas most strongly associated with linguistic functions, ______'. Câu kết luận cần khái quát hóa hai ví dụ ở câu trước.",
        "tier2_hint": "Hai ví dụ nói rằng: 1) Nếu lời nói mô tả hành động thể chất $\to$ kích hoạt vùng vận động (sensory-motor). 2) Nếu lời nói giàu cảm xúc $\to$ kích hoạt vùng cảm xúc (affective). Điều này có nghĩa là vùng não nào kích hoạt phụ thuộc vào NỘI DUNG / BỐI CẢNH của lời nói.",
        "tier3_hint": "Phương án A khái quát hoàn hảo: Sự kết hợp cụ thể của các vùng não phụ thuộc vào bối cảnh toàn diện của lời nói (contingent on the total context). Chọn A.",
        "strategic_tip": "Mẹo Princeton Review: Khi đoạn văn đưa ra 2 ví dụ cụ thể về 2 tình huống khác nhau dẫn đến 2 kết quả khác nhau, câu kết luận 'Thus' thường mang nghĩa: 'Kết quả phụ thuộc vào tình huống/bối cảnh' (contingent on context).",
        "level1_tip": "Hai ví dụ chứng minh: nội dung nói về vận động kích hoạt vùng vận động, nội dung cảm xúc kích hoạt vùng cảm xúc.",
        "level2_hint": "Khái quát hóa: Sự kết hợp các vùng não phụ thuộc vào bối cảnh cụ thể của lời nói.",
        "level3_solution": "Chọn A. the specific combination of brain regions involved in comprehending speech is likely contingent on the total context...",
        "target_clue": "Từ nối 'Thus' yêu cầu rút ra nguyên lý tổng quát từ 2 trường hợp: mô tả vận động kích hoạt vùng cảm giác - vận động; câu chuyện cảm xúc kích hoạt vùng cảm xúc.",
        "solution_steps": "1. Tiền đề: Việc hiểu ngôn ngữ không chỉ dùng vùng não ngôn ngữ mà còn kích hoạt các vùng phi ngôn ngữ.\n2. Minh chứng:\n   - Nghe về hành động cơ thể -> kích hoạt vùng cảm giác - vận động (sensory-motor).\n   - Nghe về câu chuyện xúc động -> kích hoạt vùng cảm xúc (affective).\n3. Kết luận tất yếu: Vùng não phi ngôn ngữ nào được huy động sẽ tùy thuộc vào ngữ cảnh/nội dung của phát ngôn (contingent on the total context).\n4. Đối chiếu phương án:\n   - A phản ánh chính xác quy luật phụ thuộc bối cảnh này.\n   - B sai vì nói vùng vận động kích hoạt cả khi không nói về vận động (trái với bài).\n   - C đi sâu vào suy diễn về ý định người nói (speaker did not intend).\n   - D tự ý kết luận tiêu cực 'impossible to determine' không có cơ sở.",
        "pooe_breakdown": "• Đáp án ĐÚNG: A. the specific combination of brain regions... is likely contingent on the total context... — Khái quát hóa chính xác hai ví dụ minh họa thành nguyên lý tổng quát về sự phụ thuộc vào bối cảnh phát ngôn.\n• Loại B. sensory-motor functions are involved... even in contexts that don’t involve physical movement — Sai dữ kiện; mâu thuẫn trực tiếp với phân tích trong bài.\n• Loại C. affective processing activate... even if the speaker did not intend — Đi lạc hướng sang phân tích tâm lý ý định của người nói.\n• Loại D. exact combination of areas may be impossible to determine — Thêm thắt suy đoán tiêu cực không có trong văn bản.",
        "pro_tip": "Mẹo Princeton Review: Khi đoạn văn đưa ra 2 ví dụ cụ thể về 2 tình huống khác nhau dẫn đến 2 kết quả khác nhau, câu kết luận 'Thus' thường mang nghĩa: 'Kết quả phụ thuộc vào tình huống/bối cảnh' (contingent on context).",
        "option_explanations": {
            "A": "ĐÁP ÁN ĐÚNG: A — Khái quát hóa hoàn hảo hai dẫn chứng trong bài. Việc hiểu ngôn ngữ kích hoạt vùng não nào hoàn toàn phụ thuộc vào bối cảnh và nội dung cụ thể (contingent on the total context) của lời nói đó.",
            "B": "Loại B — Mâu thuẫn với bài đọc. Tác giả chỉ rõ vùng cảm giác - vận động chỉ được gọi đến khi nghe mô tả về các hành động thể chất, không phải trong mọi ngữ cảnh chung chung.",
            "C": "Loại C — Lạc đề. Đoạn văn không thảo luận về việc người nghe suy đoán cảm xúc vượt quá ý định của người nói.",
            "D": "Loại D — Suy diễn không có căn cứ. Tác giả không hề khẳng định rằng 'không thể xác định được sự kết hợp chính xác của các vùng não'."
        }
    },
    14: {
        "id": 41,
        "master_id": "M41",
        "q_number": 14,
        "module_code": "RW_Module2",
        "topic": "Standard English Conventions (Modifiers & Pronoun Reference)",
        "difficulty": "Medium",
        "passage": "A premortem is an exercise in which teams imagine their project has already failed, then brainstorm possible reasons for that failure. Despite its apparent pessimism, ______ prompting teams to proactively develop mitigation strategies during the planning stage. As a result, the project is more likely to succeed.",
        "prompt": "Which choice completes the text so that it conforms to the conventions of Standard English?",
        "options_raw": "A. effective uncovering of potential problems by the team in this activity is done early,\nB. this activity effectively uncovers potential problems early,\nC. potential problems are effectively uncovered by this activity early,\nD. teams effectively uncover potential problems early with this activity,",
        "correct_answer": "B",
        "princeton_strategy": "Grammar - Dangling Modifier & Pronoun Antecedent",
        "tier1_hint": "Cụm mở đầu: 'Despite its apparent pessimism,' (Bất chấp sự bi quan bề ngoài của nó,). Đại từ số ít 'its' bắt buộc phải quy chiếu về một danh từ số ít chỉ sự việc: 'A premortem' (bài tập premortem) hoặc 'this activity' (hoạt động này).",
        "tier2_hint": "Chủ ngữ đứng ngay sau dấu phẩy BẮT BUỘC phải là chủ thể có vẻ bi quan nhưng mang lại hiệu quả: 'this activity'.",
        "tier3_hint": "Phương án B 'this activity effectively uncovers potential problems early,' vừa làm sáng tỏ đại từ 'its', vừa tạo cấu trúc chủ động ngắn gọn, liên kết hoàn hảo với cụm phân từ 'prompting teams...'. Chọn B.",
        "strategic_tip": "Mẹo Princeton Review: Khi câu bắt đầu bằng cụm giới từ chứa đại từ sở hữu ('its...'), chủ ngữ đứng sát sau dấu phẩy bắt buộc phải là danh từ mà 'its' thay thế!",
        "level1_tip": "Cụm mở đầu: 'Despite its apparent pessimism,' -> 'its' quy chiếu về danh từ số ít 'this activity'.",
        "level2_hint": "Loại C ('potential problems') và D ('teams') vì không ăn khớp với đại từ sở hữu số ít 'its'.",
        "level3_solution": "Chọn B. this activity effectively uncovers potential problems early,",
        "target_clue": "Đại từ sở hữu 'its' trong 'Despite its apparent pessimism' yêu cầu danh từ số ít đứng đầu mệnh đề chính phải là 'this activity' (hoạt động này).",
        "solution_steps": "1. Phân tích cụm mở đầu: 'Despite its apparent pessimism,' (Dù bề ngoài có vẻ bi quan).\n2. Đại từ 'its' là sở hữu số ít, thay thế cho 'a premortem exercise' ở câu trước. Do đó, chủ ngữ của mệnh đề chính phải là 'this activity'.\n3. Xét các phương án:\n   - A dùng cụm danh từ danh hóa rườm rà 'effective uncovering... is done early', rất vụng về.\n   - C dùng 'potential problems' (các vấn đề tiềm ẩn) làm chủ ngữ -> sai vì 'problems' không thể có 'apparent pessimism'.\n   - D dùng 'teams' (số nhiều) làm chủ ngữ -> sai vì không khớp với đại từ sở hữu số ít 'its'.\n   - B dùng 'this activity' làm chủ ngữ, kết hợp động từ chủ động 'effectively uncovers' và nối liền với 'prompting teams...'. Chọn B.",
        "pooe_breakdown": "• Đáp án ĐÚNG: B. this activity effectively uncovers potential problems early, — Chủ ngữ 'this activity' khớp hoàn hảo với đại từ 'its', tạo cấu trúc câu chủ động, mạch lạc và tự nhiên.\n• Loại A. effective uncovering of potential problems... is done early — Cấu trúc bị động danh hóa vụng về, không tự nhiên.\n• Loại C. potential problems are effectively uncovered... — Lỗi logic chủ ngữ; 'tiềm ẩn vấn đề' không thể mang tính bi quan bề ngoài ('its apparent pessimism').\n• Loại D. teams effectively uncover potential problems... — Lỗi quy chiếu đại từ; 'teams' là số nhiều mâu thuẫn với đại từ số ít 'its'.",
        "pro_tip": "Mẹo Princeton Review: Khi câu bắt đầu bằng cụm giới từ chứa đại từ sở hữu ('its...'), chủ ngữ đứng sát sau dấu phẩy bắt buộc phải là danh từ mà 'its' thay thế!",
        "option_explanations": {
            "A": "Loại A — Lối hành văn tối nghĩa và rườm rà. Việc danh hóa động từ ('effective uncovering... is done') làm câu trở nên nặng nề và mất tính liên kết với cụm phân từ phía sau.",
            "B": "ĐÁP ÁN ĐÚNG: B — Cấu trúc chuẩn xác tuyệt đối. Chủ ngữ 'this activity' (hoạt động này) liên kết mạch lạc với đại từ số ít 'its apparent pessimism', đồng thời làm chủ thể chủ động cho hành động 'prompting teams'.",
            "C": "Loại C — Lỗi sai chủ ngữ logic (Dangling Modifier). Các vấn đề tiềm ẩn ('potential problems') không phải là đối tượng có vẻ ngoài bi quan ('its apparent pessimism').",
            "D": "Loại D — Lỗi không thống nhất đại từ (Pronoun Disagreement). Chủ ngữ 'teams' là danh từ số nhiều, không thể được quy chiếu bằng đại từ sở hữu số ít 'its'."
        }
    },
    15: {
        "id": 42,
        "master_id": "M42",
        "q_number": 15,
        "module_code": "RW_Module2",
        "topic": "Standard English Conventions (Punctuation in Complex Lists)",
        "difficulty": "Hard",
        "passage": "As one of aviation’s earliest female stunt performers, Gladys Ingle developed a versatile repertoire of acrobatic feats. Her specialized skills included standing on the wing of a flying aircraft and shooting a bow and arrow at a ______ perfect balance during aerial maneuvers, such as vertical loops; and even walking across the wing as an aircraft flew at full speed beneath a bridge.",
        "prompt": "Which choice completes the text so that it conforms to the conventions of Standard English?",
        "options_raw": "A. target; maintaining\nB. target maintaining\nC. target. Maintaining\nD. target, maintaining",
        "correct_answer": "A",
        "princeton_strategy": "Grammar - Semicolons in Complex Series",
        "tier1_hint": "Quan sát cấu trúc danh sách liệt kê sau động từ 'included': 1) 'standing on the wing... and shooting a bow and arrow at a target', 2) 'maintaining perfect balance during aerial maneuvers, such as vertical loops', và 3) 'and even walking across the wing...'.",
        "tier2_hint": "Để ý mục thứ hai có chứa dấu phẩy nội bộ (', such as vertical loops;') và trước mục thứ ba có dấu chấm phẩy ('; and even walking...').",
        "tier3_hint": "Theo quy tắc ngữ pháp tiếng Anh chuẩn, khi các mục trong chuỗi liệt kê chứa dấu phẩy nội bộ, ta BẮT BUỘC phải dùng dấu chấm phẩy (;) để phân tách giữa các mục lớn. Do đó, giữa mục 1 và mục 2 phải là 'target; maintaining'. Chọn A.",
        "strategic_tip": "Mẹo Princeton Review: Dấu chấm phẩy (;) trong bài thi SAT ngoài việc nối 2 mệnh đề độc lập, còn có một nhiệm vụ đặc biệt: làm 'Super-Comma' để ngăn cách các mục trong một danh sách phức tạp đã có sẵn dấu phẩy!",
        "level1_tip": "Nhận diện chuỗi liệt kê phức tạp 3 kỹ năng có dấu chấm phẩy ở mục 3: '; and even walking...'.",
        "level2_hint": "Dùng dấu chấm phẩy (;) sau 'target' để đồng bộ với cấu trúc phân tách danh sách lớn.",
        "level3_solution": "Chọn A. target; maintaining",
        "target_clue": "Manh mối song hành: Phía sau có cấu trúc '; and even walking...' chứng minh toàn bộ chuỗi liệt kê này dùng dấu chấm phẩy làm vách ngăn cách giữa các mục.",
        "solution_steps": "1. Xác định cấu trúc câu: 'Her specialized skills included [Item 1] ; [Item 2, containing commas] ; and [Item 3]'.\n2. Item 1: 'standing on the wing... and shooting a bow and arrow at a target'.\n3. Item 2: 'maintaining perfect balance during aerial maneuvers, such as vertical loops'.\n4. Item 3: '; and even walking across the wing...'.\n5. Vì Item 2 có chứa dấu phẩy giải thích (', such as...'), và người viết đã dùng dấu chấm phẩy trước Item 3 ('; and even walking'), nên giữa Item 1 và Item 2 bắt buộc phải dùng dấu chấm phẩy: 'target; maintaining'.\n6. Loại B và D vì dùng dấu phẩy sẽ gây hỗn loạn cấu trúc liệt kê. Loại C vì dấu chấm câu sẽ biến phần sau thành câu què. Chọn A.",
        "pooe_breakdown": "• Đáp án ĐÚNG: A. target; maintaining — Sử dụng dấu chấm phẩy chuẩn xác theo quy tắc chuỗi liệt kê phức tạp (Complex Series) để đồng bộ với '; and even walking' ở cuối câu.\n• Loại B. target maintaining — Thiếu dấu ngắt câu, gây dính liền hai kỹ năng khác biệt.\n• Loại C. target. Maintaining — Dấu chấm ngắt câu biến vế sau thành mảnh câu què (Sentence Fragment) vì không có vị ngữ chính.\n• Loại D. target, maintaining — Dấu phẩy không đủ cấp độ để phân tách các mục lớn đã chứa dấu phẩy nội bộ.",
        "pro_tip": "Mẹo Princeton Review: Dấu chấm phẩy (;) trong bài thi SAT ngoài việc nối 2 mệnh đề độc lập, còn có một nhiệm vụ đặc biệt: làm 'Super-Comma' để ngăn cách các mục trong một danh sách phức tạp đã có sẵn dấu phẩy!",
        "option_explanations": {
            "A": "ĐÁP ÁN ĐÚNG: A (target; maintaining) — Chuẩn xác quy tắc chuỗi liệt kê phức tạp (Complex Series). Khi các thành phần trong chuỗi liệt kê đã có dấu phẩy nội bộ (như ', such as vertical loops'), dấu chấm phẩy (;) được dùng làm dấu phân cách chính giữa các mục lớn, đồng bộ với '; and even walking' ở vế sau.",
            "B": "Loại B — Lỗi chạy chữ không dấu câu (Run-on). Hai kỹ năng riêng biệt bị ghép dính vào nhau mà không có bất kỳ dấu hiệu phân cách nào.",
            "C": "Loại C — Lỗi câu chưa hoàn chỉnh (Sentence Fragment). Dấu chấm kết thúc câu tại 'target.' sẽ biến toàn bộ phần còn lại phía sau thành một đoạn cụm từ không có động từ chính.",
            "D": "Loại D — Dấu phẩy không đủ thẩm quyền phân tách trong danh sách phức tạp, gây nhầm lẫn cấp độ với dấu phẩy giải thích ', such as vertical loops'."
        }
    },
    16: {
        "id": 43,
        "master_id": "M43",
        "q_number": 16,
        "module_code": "RW_Module2",
        "topic": "Standard English Conventions (Finite vs Nonfinite Verbs)",
        "difficulty": "Medium",
        "passage": "During a recent investigation, archaeologists ______ an explanation for the consistent and seemingly deliberate placement of petroglyphs (rock carvings) near water analyzed several sites on Caribbean islands. They discovered that many of the pools near petroglyphs had consistent water levels year-round, creating perfect reflecting surfaces for the carved symbols on the rocks.",
        "prompt": "Which choice completes the text so that it conforms to the conventions of Standard English?",
        "options_raw": "A. had sought\nB. sought\nC. seeking\nD. were seeking",
        "correct_answer": "C",
        "princeton_strategy": "Grammar - Verbs vs Modifiers (Main Verb Identification)",
        "tier1_hint": "Quét toàn bộ câu để tìm động từ chính đã chia thì: Chủ ngữ là 'archaeologists' (các nhà khảo cổ) và động từ chính vị ngữ là 'analyzed' (đã phân tích vài địa điểm trên các đảo Caribbean).",
        "tier2_hint": "Vì câu ĐÃ CÓ động từ chính chia thì là 'analyzed', phần đứng giữa chủ ngữ và động từ chính bắt buộc phải là một cụm phân từ bổ nghĩa (participial modifier), KHÔNG ĐƯỢC là động từ chia thì thứ hai.",
        "tier3_hint": "Trong 4 phương án, chỉ có 'seeking' (hiện tại phân từ V-ing) là dạng bổ nghĩa không chia thì, bổ nghĩa cho 'archaeologists'. Các phương án A, B, D đều là động từ chia thì sẽ tạo thành lỗi thừa vị ngữ (Double Verb). Chọn C.",
        "strategic_tip": "Mẹo Princeton Review: Luôn tìm động từ chính (Main Verb) của câu trước! Nếu câu đã có một động từ chính hoàn chỉnh như 'analyzed', chỗ trống chỉ có thể là dạng phân từ bổ nghĩa (-ing hoặc -ed).",
        "level1_tip": "Tìm động từ chính của câu: 'analyzed several sites on Caribbean islands'.",
        "level2_hint": "Vì đã có động từ chính 'analyzed', chỗ trống phải là dạng phân từ bổ nghĩa 'seeking' (-ing).",
        "level3_solution": "Chọn C. seeking (đang tìm kiếm lời giải thích).",
        "target_clue": "Động từ chính của câu là 'analyzed'. Cụm 'seeking an explanation...' là phân từ hiện tại bổ nghĩa cho chủ ngữ 'archaeologists'.",
        "solution_steps": "1. Phân tích cấu trúc câu:\n   - Trạng ngữ: 'During a recent investigation,'\n   - Chủ ngữ: 'archaeologists'\n   - Thành phần bổ nghĩa: '[seeking an explanation for... near water]'\n   - Vị ngữ chính: 'analyzed several sites on Caribbean islands.'\n2. Nếu điền A (had sought), B (sought), hay D (were seeking), câu sẽ có 2 động từ chia thì đứng liền nhau ('archaeologists sought ... analyzed ...') mà không có liên từ nối, tạo thành câu sai cấu trúc ngữ pháp nghiêm trọng.\n3. Do đó, chỗ trống phải dùng dạng phân từ phi vị ngữ 'seeking' để bổ nghĩa cho 'archaeologists'. Chọn C.",
        "pooe_breakdown": "• Đáp án ĐÚNG: C. seeking — Hiện tại phân từ đóng vai trò tính từ bổ nghĩa cho chủ ngữ 'archaeologists', cho phép 'analyzed' đóng vai trò động từ chính duy nhất của câu.\n• Loại A. had sought — Động từ chia thì quá khứ hoàn thành; tạo lỗi câu có 2 vị ngữ chính thừa thãi.\n• Loại B. sought — Động từ chia thì quá khứ đơn; tạo lỗi câu hai vị ngữ không liên từ.\n• Loại D. were seeking — Động từ chia thì quá khứ tiếp diễn; tạo lỗi thừa động từ chia thì.",
        "pro_tip": "Mẹo Princeton Review: Luôn tìm động từ chính (Main Verb) của câu trước! Nếu câu đã có một động từ chính hoàn chỉnh như 'analyzed', chỗ trống chỉ có thể là dạng phân từ bổ nghĩa (-ing hoặc -ed).",
        "option_explanations": {
            "A": "Loại A — Lỗi thừa vị ngữ chính (Double Finite Verb). Nếu dùng 'had sought', câu sẽ có hai động từ chia thì cùng tranh giành một chủ ngữ mà không có liên từ kết nối ('archaeologists had sought ... analyzed ...').",
            "B": "Loại B — Lỗi hai động từ vị ngữ đứng liền. Động từ quá khứ đơn 'sought' xung đột ngữ pháp với động từ chính 'analyzed'.",
            "C": "ĐÁP ÁN ĐÚNG: C (seeking) — Chuẩn xác ngữ pháp. 'seeking' là hiện tại phân từ (present participle) tạo thành cụm phân từ bổ nghĩa cho 'archaeologists' (các nhà khảo cổ, những người đang tìm kiếm lời giải thích...), giữ cho 'analyzed' là động từ chính duy nhất của câu.",
            "D": "Loại D — Lỗi vị ngữ chia thì tiếp diễn. 'were seeking' không thể kết hợp hợp lệ với động từ chính 'analyzed' đứng sau."
        }
    },
    17: {
        "id": 44,
        "master_id": "M44",
        "q_number": 17,
        "module_code": "RW_Module2",
        "topic": "Standard English Conventions (Punctuation & Parenthetical Phrases)",
        "difficulty": "Hard",
        "passage": "For his mash-up “Smells Like Rockin’ Robin,” DJ Mark Vidler layered Kurt Cobain’s minor-key vocals from “Smells Like Teen Spirit” over the Jackson 5’s major-key blues progression from “Rockin’ ______ by inserting the grunge singer’s vocal melody as pitch-shifted “blue notes” within the pop-soul group’s harmonic framework, creates an intensified “bluesy” sound and subverts the expectations of listeners familiar with either musical source.",
        "prompt": "Which choice completes the text so that it conforms to the conventions of Standard English?",
        "options_raw": "A. Robin,” a technique that\nB. Robin.” A technique that\nC. Robin”; a technique that,\nD. Robin”—a technique that,",
        "correct_answer": "D",
        "princeton_strategy": "Grammar - Em Dash & Nonessential Modifiers",
        "tier1_hint": "Nhìn cấu trúc tổng thể: Mệnh đề chính 'DJ Mark Vidler layered... over the Jackson 5’s blues progression from Rockin' Robin'. Phía sau là cụm đồng vị ngữ mở rộng giải thích cho kỹ thuật này: 'a technique that... creates an intensified bluesy sound...'.",
        "tier2_hint": "Lưu ý cụm chèn ở giữa: 'by inserting the grunge singer’s vocal melody... framework' là một cụm trạng từ xen giữa có dấu phẩy đóng ở cuối (', creates an...'). Do đó, trước cụm này bắt buộc phải có một dấu phẩy mở: ', by inserting...'.",
        "tier3_hint": "Để kết nối câu chính với cụm đồng vị ngữ mở rộng 'a technique that, by inserting..., creates...', tác giả sử dụng dấu gạch ngang (em dash): 'Robin\"—a technique that,'. Chọn D.",
        "strategic_tip": "Mẹo Princeton Review: Dấu gạch ngang (em dash) là công cụ hoàn hảo trong SAT để gắn một cụm đồng vị ngữ mở rộng (Appositive Modifier) vào cuối mệnh đề độc lập mà không gây nhầm lẫn với các dấu phẩy nội bộ.",
        "level1_tip": "Cần dấu gạch ngang để tách cụm giải thích 'a technique that...' và dấu phẩy mở cho cụm xen giữa ', by inserting...'.",
        "level2_hint": "Loại A vì thiếu dấu phẩy mở cho cụm chèn, loại B vì tạo câu què, loại C vì chấm phẩy sai quy tắc.",
        "level3_solution": "Chọn D. Robin”—a technique that,",
        "target_clue": "Cặp dấu phẩy cho cụm xen giữa: ', by inserting... harmonic framework,' đòi hỏi phải có dấu phẩy sau 'that'. Dấu gạch ngang (em dash) gắn kết hoàn hảo cụm đồng vị ngữ 'a technique that...'.",
        "solution_steps": "1. Mệnh đề 1 hoàn chỉnh: 'DJ Mark Vidler layered Kurt Cobain’s vocals... over the blues progression from Rockin' Robin'.\n2. Phần sau bổ nghĩa: 'a technique that [by inserting...,] creates an intensified sound...'.\n3. Cụm 'by inserting...' là trạng ngữ xen giữa, đã có dấu phẩy đóng trước 'creates'. Do đó bắt buộc phải có dấu phẩy mở ngay sau 'that': 'that, by inserting...'.\n4. Để liên kết mệnh đề độc lập 1 với cụm đồng vị ngữ 'a technique that,', dấu gạch ngang (em dash) là dấu câu chính xác và ấn tượng nhất.\n5. Loại A vì thiếu dấu phẩy sau 'that'. Loại B vì dấu chấm biến phần sau thành câu què. Loại C vì chấm phẩy không thể đứng trước cụm danh từ phụ thuộc. Chọn D.",
        "pooe_breakdown": "• Đáp án ĐÚNG: D. Robin”—a technique that, — Dấu gạch ngang mở cụm đồng vị ngữ bổ nghĩa, dấu phẩy sau 'that' mở đầu chính xác cụm xen giữa ', by inserting...'.\n• Loại A. Robin,” a technique that — Thiếu dấu phẩy mở cho cụm trạng ngữ chèn giữa ', by inserting...'.\n• Loại B. Robin.” A technique that — Tạo thành mảnh câu què (Sentence Fragment) vì vế sau chỉ là mệnh đề quan hệ.\n• Loại C. Robin”; a technique that, — Dấu chấm phẩy (;) không được dùng trước một cụm danh từ bổ nghĩa.",
        "pro_tip": "Mẹo Princeton Review: Dấu gạch ngang (em dash) là công cụ hoàn hảo trong SAT để gắn một cụm đồng vị ngữ mở rộng (Appositive Modifier) vào cuối mệnh đề độc lập mà không gây nhầm lẫn với các dấu phẩy nội bộ.",
        "option_explanations": {
            "A": "Loại A — Thiếu dấu câu cho thành phần chèn giữa (Missing Comma). Cụm trạng từ 'by inserting...' có dấu phẩy ở đuôi (', creates') nhưng lại thiếu dấu phẩy mở đầu ngay sau 'that'.",
            "B": "Loại B — Lỗi câu chưa hoàn chỉnh (Sentence Fragment). Dấu chấm câu biến phần 'A technique that...' thành một câu què không có chủ - vị độc lập.",
            "C": "Loại C — Lỗi chấm phẩy sai quy tắc. Dấu chấm phẩy bắt buộc phải nối hai mệnh đề độc lập hoàn chỉnh, không thể dùng trước cụm bổ nghĩa phụ thuộc.",
            "D": "ĐÁP ÁN ĐÚNG: D (Robin”—a technique that,) — Chuẩn xác ngữ pháp. Dấu gạch ngang (em dash) tạo điểm nhấn liên kết cụm đồng vị ngữ mở rộng với mệnh đề trước, và dấu phẩy sau 'that' tạo thành cặp dấu phẩy hoàn chỉnh đóng khung cụm trạng ngữ xen giữa ', by inserting...'."
        }
    }
}

def update_json_file(file_path):
    with open(file_path, 'r', encoding='utf-8') as f:
        data = json.load(f)
    
    count = 0
    for q in data.get('questions', []):
        q_num = q.get('q_number')
        if q.get('module_code') == 'RW_Module2' and q_num in AUDIT_FIXES:
            info = AUDIT_FIXES[q_num]
            for k, v in info.items():
                q[k] = v
            count += 1
            print(f"[{file_path}] Updated Q{q_num} (ID {q.get('id')})")
            
    with open(file_path, 'w', encoding='utf-8') as f:
        json.dump(data, f, ensure_ascii=False, indent=2)
    print(f"Successfully saved {count} questions to {file_path}")

def update_excel_explanations(file_path):
    wb = openpyxl.load_workbook(file_path)
    
    # 1. Master Sheet
    if 'Master_All_98_Explanations' in wb.sheetnames:
        ws = wb['Master_All_98_Explanations']
        for r in range(5, 103):
            mid = ws.cell(r, 1).value
            mod = ws.cell(r, 2).value
            q_val = ws.cell(r, 3).value
            for q_num, info in AUDIT_FIXES.items():
                if mid == info["master_id"] or (mod == 'RW_Module2' and q_val in [q_num, str(q_num), f"Q{q_num}"]):
                    ws.cell(r, 4).value = info["topic"]
                    ws.cell(r, 5).value = info["difficulty"]
                    ws.cell(r, 6).value = info["passage"]
                    ws.cell(r, 7).value = info["prompt"]
                    ws.cell(r, 8).value = info["options_raw"]
                    ws.cell(r, 9).value = info["correct_answer"]
                    ws.cell(r, 10).value = info["princeton_strategy"]
                    ws.cell(r, 11).value = info["target_clue"]
                    ws.cell(r, 12).value = info["solution_steps"]
                    ws.cell(r, 13).value = info["pooe_breakdown"]
                    ws.cell(r, 14).value = info["pro_tip"]
                    print(f"Updated Master_All_98_Explanations row {r} for {mid}")

    # 2. RW_Module2_Explanations Sheet
    if 'RW_Module2_Explanations' in wb.sheetnames:
        ws = wb['RW_Module2_Explanations']
        for r in range(5, ws.max_row + 1):
            q_val = ws.cell(r, 1).value
            for q_num, info in AUDIT_FIXES.items():
                if q_val in [q_num, str(q_num), f"Q{q_num}"]:
                    ws.cell(r, 2).value = info["correct_answer"]
                    ws.cell(r, 3).value = info["target_clue"]
                    ws.cell(r, 4).value = info["solution_steps"]
                    ws.cell(r, 5).value = info["pooe_breakdown"]
                    ws.cell(r, 6).value = info["topic"]
                    ws.cell(r, 7).value = info["difficulty"]
                    ws.cell(r, 8).value = info["passage"]
                    ws.cell(r, 9).value = info["prompt"]
                    ws.cell(r, 10).value = info["options_raw"]
                    ws.cell(r, 11).value = info["princeton_strategy"]
                    ws.cell(r, 12).value = info["pro_tip"]
                    print(f"Updated RW_Module2_Explanations row {r} for Q{q_num}")

    wb.save(file_path)
    print(f"Successfully saved {file_path}")

def update_excel_template(file_path):
    wb = openpyxl.load_workbook(file_path)
    if 'RW_Module2' in wb.sheetnames:
        ws = wb['RW_Module2']
        for r in range(1, ws.max_row + 1):
            q_val = ws.cell(r, 1).value
            for q_num, info in AUDIT_FIXES.items():
                if q_val in [q_num, str(q_num), f"Q{q_num}"]:
                    ws.cell(r, 2).value = info["topic"]
                    ws.cell(r, 3).value = info["passage"]
                    ws.cell(r, 4).value = info["prompt"]
                    ws.cell(r, 5).value = info["options_raw"]
                    print(f"Updated {file_path} sheet RW_Module2 row {r} for Q{q_num}")
                    
    if 'Master_All_98_Questions' in wb.sheetnames:
        ws = wb['Master_All_98_Questions']
        for r in range(1, ws.max_row + 1):
            mid = ws.cell(r, 1).value
            for q_num, info in AUDIT_FIXES.items():
                if mid == info["master_id"]:
                    ws.cell(r, 3).value = info["topic"]
                    ws.cell(r, 4).value = info["passage"]
                    ws.cell(r, 5).value = info["prompt"]
                    ws.cell(r, 6).value = info["options_raw"]
                    print(f"Updated {file_path} sheet Master_All_98_Questions row {r} for {mid}")
                    
    wb.save(file_path)
    print(f"Successfully saved {file_path}")

if __name__ == '__main__':
    update_json_file('SAT Coach/BluebookOct2026App/database.json')
    update_json_file('SAT Coach/BluebookOct2026App/static/database.json')
    update_excel_explanations('SAT Coach/Real Tests/Oct 2026/Explanations.xlsx')
    update_excel_template('SAT Coach/BluebookOct2026App/template.xlsx')
    update_excel_template('SAT Coach/BluebookOct2026App/static/template.xlsx')
