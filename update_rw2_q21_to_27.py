import json
import openpyxl

NEW_DATA_RW2 = {
    21: {
        "id": 48,
        "master_id": "M48",
        "q_number": 21,
        "module_code": "RW_Module2",
        "topic": "Transitions & Discourse Connectors (Contrast)",
        "difficulty": "Medium",
        "passage": "In a 1993 study by Vales and Peek, the researchers used microhistological fecal analysis to determine the ratio of three plant subtypes (graminoids, forbs, and browse) within the diets of North American ungulates. The researchers did not perform this analysis on all such ungulates, ______ they focused exclusively on elk in Wyoming.",
        "prompt": "Which choice completes the text with the most logical transition?",
        "options_raw": "A. for instance;\nB. in other words;\nC. regardless;\nD. however;",
        "correct_answer": "D",
        "princeton_strategy": "Transitions - Contrast & Scope Restriction",
        "tier1_hint": "Xác định mối quan hệ logic giữa hai vế câu: Vế trước nói về việc 'không thực hiện trên tất cả các loài móng guốc' (did not perform on all), vế sau giới hạn phạm vi 'chỉ tập trung duy nhất vào loài nai sừng tấm ở Wyoming' (focused exclusively on elk).",
        "tier2_hint": "Đây là quan hệ tương phản / đối lập giới hạn phạm vi nghiên cứu (Contrast). Hãy loại các từ nối mang ý nghĩa minh họa (for instance) hoặc diễn giải tương đương (in other words).",
        "tier3_hint": "Từ nối 'however;' thể hiện sự tương phản rõ rệt: dù không nghiên cứu toàn bộ, tuy nhiên họ tập trung sâu vào một đối tượng cụ thể là nai sừng tấm Wyoming. Chọn D.",
        "strategic_tip": "Mẹo Princeton Review: Khi vế trước dùng phủ định (did not / never) và vế sau thu hẹp phạm vi vào một đối tượng cá biệt, từ nối đối lập 'however;' là lựa chọn tối ưu.",
        "level1_tip": "Nhận diện quan hệ tương phản giữa việc 'không làm trên tất cả' và 'chỉ tập trung vào một loài'.",
        "level2_hint": "Loại bỏ từ nối ví dụ (A) và diễn giải lại (B).",
        "level3_solution": "Chọn D. however; (tuy nhiên).",
        "target_clue": "Mệnh đề 1: 'did not perform this analysis on all such ungulates' (phủ định toàn thể) đối lập với Mệnh đề 2: 'they focused exclusively on elk' (khẳng định thu hẹp).",
        "solution_steps": "1. Phân tích ngữ cảnh: Các nhà nghiên cứu không phân tích trên toàn bộ các loài động vật móng guốc Bắc Mỹ, TUY NHIÊN (however) họ chỉ tập trung phân tích mẫu phân của loài nai sừng tấm (elk) tại Wyoming.\n2. Quan hệ logic là sự tương phản đối lập (Contrast) nhằm thu hẹp phạm vi khảo sát.\n3. Loại A (for instance) vì nai sừng tấm không phải là ví dụ cho việc 'không làm trên tất cả'. Loại B (in other words) vì câu sau không nói lại cùng nghĩa. Loại C (regardless) dùng sai ngữ nghĩa liên kết câu. Chọn D.",
        "pooe_breakdown": "• Đáp án ĐÚNG: D. however; — Thể hiện chính xác mối quan hệ tương phản, đối lập logic giữa việc không nghiên cứu diện rộng và việc tập trung cục bộ vào nai sừng tấm.\n• Loại A. for instance; — Sai logic ví dụ; vế sau không phải ví dụ minh họa cho mệnh đề phủ định trước.\n• Loại B. in other words; — Sai logic diễn giải; hai câu mang hai nội dung thông tin khác biệt.\n• Loại C. regardless; — 'Bất kể / Dù sao đi nữa' không phản ánh đúng quan hệ giới hạn nghiên cứu.",
        "pro_tip": "Mẹo Princeton Review: Khi vế trước dùng phủ định (did not / never) và vế sau thu hẹp phạm vi vào một đối tượng cá biệt, từ nối đối lập 'however;' là lựa chọn tối ưu.",
        "option_explanations": {
            "A": "Loại A (for instance;) — Sai quan hệ logic. Mệnh đề sau không đưa ra ví dụ minh họa cho việc 'không làm trên tất cả', mà thể hiện sự tương phản đối lập về phạm vi nghiên cứu.",
            "B": "Loại B (in other words;) — Sai logic diễn giải. Câu sau không nhắc lại cùng một ý bằng từ ngữ khác, mà giới thiệu đối tượng nghiên cứu cụ thể.",
            "C": "Loại C (regardless;) — 'Bất kể / Bất chấp' không diễn đạt được sự chuyển tiếp logic giữa mệnh đề phủ định trước và khẳng định sau.",
            "D": "ĐÁP ÁN ĐÚNG: D (however;) — Thể hiện chính xác mối quan hệ tương phản (contrast), chuyển tiếp từ thực tế 'không nghiên cứu toàn bộ' sang hành động 'tuy nhiên chỉ tập trung vào một loài cụ thể'."
        }
    },
    22: {
        "id": 49,
        "master_id": "M49",
        "q_number": 22,
        "module_code": "RW_Module2",
        "topic": "Transitions & Discourse Connectors (Cause & Effect / Conclusion)",
        "difficulty": "Medium",
        "passage": "The Hornbostel-Sachs system classifies musical instruments by how they produce sound. For example, an instrument that is popular in India called the saraswati veena produces sound primarily through the vibration of its strings. ______ under the Hornbostel-Sachs system, the saraswati veena is a chordophone.",
        "prompt": "Which choice completes the text with the most logical transition?",
        "options_raw": "A. Furthermore,\nB. Thus,\nC. Similarly,\nD. Instead,",
        "correct_answer": "B",
        "princeton_strategy": "Transitions - Cause, Effect & Deductive Conclusion",
        "tier1_hint": "Xác định mối quan hệ giữa câu 2 và câu 3: Câu 2 giải thích đàn saraswati veena tạo âm bằng cách rung dây. Câu 3 kết luận rằng theo hệ thống Hornbostel-Sachs, đàn này là một 'chordophone' (nhạc cụ dây).",
        "tier2_hint": "Đây là mối quan hệ nhân quả - kết luận suy diễn (Cause and Effect / Conclusion). Vì phát ra âm thanh từ dây đàn, NÊN (Thus / Therefore) nó được xếp vào nhóm chordophone.",
        "tier3_hint": "Từ nối 'Thus,' biểu thị kết quả suy luận tất yếu từ tiền đề đã nêu ở câu trước. Chọn B.",
        "strategic_tip": "Mẹo Princeton Review: Khi câu trước đưa ra cơ chế/nguyên lý cụ thể và câu sau rút ra kết luận xếp loại, hãy tìm từ nối mang ý nghĩa 'Do đó, Vì vậy' (Thus, Therefore, Hence).",
        "level1_tip": "Nhận diện quan hệ nhân quả: phát ra âm nhờ dây rung -> do đó được xếp là chordophone.",
        "level2_hint": "Loại bỏ từ nối bổ sung (A), so sánh tương đồng (C) và thay thế (D).",
        "level3_solution": "Chọn B. Thus, (Do đó / Vì vậy).",
        "target_clue": "Tiền đề: 'produces sound primarily through the vibration of its strings' -> Kết luận phân loại: 'the saraswati veena is a chordophone'.",
        "solution_steps": "1. Câu trước: Nhạc cụ saraswati veena tạo âm thanh chủ yếu qua sự rung động của các sợi dây đàn.\n2. Câu sau: Đàn này được định danh là 'chordophone' (nhạc cụ bộ dây) theo hệ thống Hornbostel-Sachs.\n3. Quan hệ là suy luận hệ quả tất yếu (Conclusion): Bởi vì rung dây nên suy ra nó là chordophone. Do đó, cần từ nối chỉ kết quả 'Thus' (Vì vậy, Do đó).\n4. Loại A (Furthermore) vì không phải bổ sung thông tin mới độc lập. Loại C (Similarly) vì không có nhạc cụ thứ hai để so sánh. Loại D (Instead) vì không có sự thay thế phủ định. Chọn B.",
        "pooe_breakdown": "• Đáp án ĐÚNG: B. Thus, — Nối kết quả suy luận logic từ tiền đề đàn phát ra âm bằng dây sang kết luận phân loại chordophone.\n• Loại A. Furthermore, — Sai logic; 'Hơn nữa' dùng để bổ sung luận điểm mới, không biểu thị kết luận nhân quả.\n• Loại C. Similarly, — Sai logic; 'Tương tự' đòi hỏi phải có một nhạc cụ thứ hai được so sánh song song.\n• Loại D. Instead, — Sai logic; 'Thay vào đó' dùng để biểu thị sự thay thế đối lập sau một mệnh đề phủ định.",
        "pro_tip": "Mẹo Princeton Review: Khi câu trước đưa ra cơ chế/nguyên lý cụ thể và câu sau rút ra kết luận xếp loại, hãy tìm từ nối mang ý nghĩa 'Do đó, Vì vậy' (Thus, Therefore, Hence).",
        "option_explanations": {
            "A": "Loại A (Furthermore,) — Sai logic bổ sung. 'Furthermore' dùng để đưa thêm thông tin phụ trợ mới, không thể hiện kết luận suy diễn tất yếu từ câu trước.",
            "B": "ĐÁP ÁN ĐÚNG: B (Thus,) — Chuẩn xác quan hệ nhân quả/suy luận. Vì đàn veena phát ra âm thanh từ độ rung của dây, DO ĐÓ (Thus) theo hệ thống Hornbostel-Sachs nó được phân loại là chordophone.",
            "C": "Loại C (Similarly,) — Sai logic so sánh. 'Similarly' đòi hỏi một đối tượng tương đồng thứ hai, trong khi đoạn văn chỉ đang phân tích duy nhất đàn veena.",
            "D": "Loại D (Instead,) — Sai logic thay thế. 'Instead' chỉ dùng khi có sự phủ định hoặc thay thế một phương án khác."
        }
    },
    23: {
        "id": 50,
        "master_id": "M50",
        "q_number": 23,
        "module_code": "RW_Module2",
        "topic": "Transitions & Discourse Connectors (Exemplification)",
        "difficulty": "Medium",
        "passage": "Generally, sleek vehicles are more aerodynamic than bulkier ones. The streamlined nose of the Airbus A350 XWB jet, ______ helps it glide through wind with relative ease, while a boxy pickup truck encounters more wind resistance, making it less aerodynamic.",
        "prompt": "Which choice completes the text with the most logical transition?",
        "options_raw": "A. additionally,\nB. however,\nC. for example,\nD. meanwhile,",
        "correct_answer": "C",
        "princeton_strategy": "Transitions - Exemplification & General to Specific",
        "tier1_hint": "So sánh câu 1 và câu 2: Câu 1 nêu nguyên lý tổng quát (sleek vehicles are more aerodynamic than bulkier ones). Câu 2 đưa ra trường hợp cụ thể: máy bay Airbus mũi thon gọn (sleek) so với xe bán tải hình hộp vuông vức (bulky).",
        "tier2_hint": "Mối quan hệ đi từ nguyên lý khái quát đến trường hợp thực tế điển hình chính là quan hệ ví dụ minh họa (Exemplification).",
        "tier3_hint": "Từ nối 'for example,' đặt giữa cặp dấu phẩy là liên từ chuẩn xác để giới thiệu trường hợp máy bay Airbus làm ví dụ cho nguyên lý khí động học. Chọn C.",
        "strategic_tip": "Mẹo Princeton Review: Chuyển dịch từ quy luật chung ('Generally...') sang đối tượng cụ thể có tên riêng ('Airbus A350 XWB jet...') luôn báo hiệu mối quan hệ Đưa ra ví dụ ('for example', 'for instance').",
        "level1_tip": "Quy luật chung (Generally) được minh họa bằng trường hợp cụ thể (máy bay Airbus vs xe bán tải).",
        "level2_hint": "Loại từ nối bổ sung (A), tương phản (B) và đồng thời (D).",
        "level3_solution": "Chọn C. for example, (ví dụ như).",
        "target_clue": "Nguyên tắc chung: 'Generally, sleek vehicles are more aerodynamic...' -> Minh chứng thực tế: 'The streamlined nose of the Airbus A350 XWB jet...'.",
        "solution_steps": "1. Câu 1: Nhìn chung, các phương tiện thon gọn thì có tính khí động học tốt hơn các phương tiện cồng kềnh.\n2. Câu 2: Mũi thuôn của máy bay Airbus A350 lướt gió dễ dàng, trong khi xe bán tải dạng hộp cản gió nhiều hơn.\n3. Đây là ví dụ trực quan chứng minh cho quy luật khái quát ở câu đầu. Cần từ nối chỉ ví dụ: 'for example,'.\n4. Loại A (additionally) vì câu 2 không thêm ý mới mà minh họa cho câu 1. Loại B (however) vì máy bay lướt gió tốt đồng nghĩa với việc thon gọn khí động học, không hề đối lập. Loại D (meanwhile) sai ngữ nghĩa. Chọn C.",
        "pooe_breakdown": "• Đáp án ĐÚNG: C. for example, — Đóng vai trò liên từ trạng ngữ minh họa một trường hợp thực tế điển hình cho nguyên lý khái quát ở câu trước.\n• Loại A. additionally, — Sai logic; câu 2 không bổ sung luận điểm độc lập mà là bằng chứng minh họa cho câu 1.\n• Loại B. however, — Sai logic; không có sự mâu thuẫn hay tương phản nào với quy luật ở câu 1.\n• Loại D. meanwhile, — Sai logic; 'trong khi đó' chỉ hai diễn biến song hành về thời gian, không dùng để nêu ví dụ.",
        "pro_tip": "Mẹo Princeton Review: Chuyển dịch từ quy luật chung ('Generally...') sang đối tượng cụ thể có tên riêng ('Airbus A350 XWB jet...') luôn báo hiệu mối quan hệ Đưa ra ví dụ ('for example', 'for instance').",
        "option_explanations": {
            "A": "Loại A (additionally,) — Sai quan hệ logic. 'Additionally' dùng để bổ sung một ý hoàn toàn mới, trong khi câu này đang lấy ví dụ thực tiễn để làm sáng tỏ câu trước.",
            "B": "Loại B (however,) — Sai logic tương phản. Việc máy bay Airbus thon gọn lướt gió tốt hoàn toàn nhất quán với nhận định 'sleek vehicles are more aerodynamic', không hề có sự đối lập.",
            "C": "ĐÁP ÁN ĐÚNG: C (for example,) — Chuẩn xác quan hệ ví dụ minh họa. Máy bay Airbus và xe bán tải là ví dụ điển hình chứng minh cho quy luật khí động học tổng quát ở câu đầu.",
            "D": "Loại D (meanwhile,) — Sai logic thời gian. 'Meanwhile' biểu thị hành động xảy ra cùng lúc, không dùng để kết nối giữa nguyên lý và ví dụ."
        }
    },
    24: {
        "id": 51,
        "master_id": "M51",
        "q_number": 24,
        "module_code": "RW_Module2",
        "topic": "Transitions & Discourse Connectors (Emphasis / Reinforcement)",
        "difficulty": "Hard",
        "passage": "On June 13, 1927, aviator Charles Lindbergh received a ticker-tape parade in New York City in recognition of his transatlantic flight. Of the 206 ticker-tape parades held between 1886 and 2022, a number were for achievements in aviation. ______ Lindbergh’s parade was just one of 16 honoring pilots.",
        "prompt": "Which choice completes the text with the most logical transition?",
        "options_raw": "A. Therefore,\nB. Indeed,\nC. In addition,\nD. Nevertheless,",
        "correct_answer": "B",
        "princeton_strategy": "Transitions - Emphasis, Elaboration & Factual Confirmation",
        "tier1_hint": "Xem xét mối liên hệ giữa câu 2 và câu 3: Câu 2 nhận định khái quát rằng 'một số cuộc diễu hành dành cho thành tựu hàng không' (a number were for achievements in aviation). Câu 3 cung cấp số liệu thực tế cụ thể khẳng định điều đó: cuộc diễu hành của Lindbergh là một trong 16 cuộc vinh danh phi công.",
        "tier2_hint": "Khi câu sau cung cấp chi tiết xác thực, khẳng định và củng cố mạnh mẽ nhận định ở câu trước, liên từ nhấn mạnh 'Indeed,' (Thật vậy / Thực tế là) là từ nối chính xác.",
        "tier3_hint": "Từ nối 'Indeed,' củng cố sự thật cho cụm 'a number were for achievements in aviation' bằng con số cụ thể 'one of 16 honoring pilots'. Chọn B.",
        "strategic_tip": "Mẹo Princeton Review: Từ nối 'Indeed' hoặc 'In fact' được bài thi SAT sử dụng khi câu sau xác nhận hoặc đào sâu mức độ cụ thể của khẳng định ở câu trước.",
        "level1_tip": "Câu 2 nói 'có một số cuộc diễu hành hàng không' -> Câu 3 khẳng định số liệu cụ thể: 'một trong 16 cuộc'.",
        "level2_hint": "Loại từ nối nguyên nhân kết quả (A), thêm vào (C) và tương phản (D).",
        "level3_solution": "Chọn B. Indeed, (Thật vậy / Thực tế là).",
        "target_clue": "Câu trước: 'a number were for achievements in aviation' -> Câu sau củng cố dữ liệu thực tế: 'Lindbergh's parade was just one of 16 honoring pilots'.",
        "solution_steps": "1. Câu 2: Trong số 206 cuộc diễu hành, có một số cuộc nhằm tôn vinh thành tựu hàng không.\n2. Câu 3: Khẳng định con số thực tế: Lindbergh chỉ là một trong 16 phi công được vinh danh.\n3. Mối quan hệ là nhấn mạnh và củng cố sự thật (Emphasis / Confirmation) cho nhận định trước. Từ nối chuẩn là 'Indeed,' (Thực vậy / Quả thực).\n4. Loại A (Therefore) vì câu 3 không phải là kết quả nhân quả sinh ra từ câu 2. Loại C (In addition) vì không thêm sự kiện mới mà làm rõ câu 2. Loại D (Nevertheless) vì không có sự đối lập. Chọn B.",
        "pooe_breakdown": "• Đáp án ĐÚNG: B. Indeed, — Dùng để nhấn mạnh và củng cố tính xác thực cho nhận định 'a number were for achievements in aviation' bằng con số cụ thể 16 cuộc diễu hành.\n• Loại A. Therefore, — Sai logic nhân quả; việc có 16 cuộc diễu hành không phải hệ quả do câu trước tạo ra.\n• Loại C. In addition, — Sai logic; câu 3 không bổ sung một sự kiện mới ngoài lề mà đào sâu chi tiết cho chính câu 2.\n• Loại D. Nevertheless, — Sai logic tương phản; không có sự đối lập hay nhượng bộ giữa hai câu.",
        "pro_tip": "Mẹo Princeton Review: Từ nối 'Indeed' hoặc 'In fact' được bài thi SAT sử dụng khi câu sau xác nhận hoặc đào sâu mức độ cụ thể của khẳng định ở câu trước.",
        "option_explanations": {
            "A": "Loại A (Therefore,) — Sai logic nhân quả. Việc Lindbergh là 1 trong 16 phi công không phải là kết quả tất yếu sinh ra từ việc có 206 cuộc diễu hành.",
            "B": "ĐÁP ÁN ĐÚNG: B (Indeed,) — Chuẩn xác ngữ nghĩa nhấn mạnh (Emphasis/Confirmation). 'Indeed' được dùng để khẳng định và bổ sung con số thực tế (1 trong 16 cuộc) củng cố cho nhận định 'a number' ở câu trước.",
            "C": "Loại C (In addition,) — Sai logic bổ sung. Câu 3 làm rõ cho nội dung câu 2 chứ không phải đưa thêm một sự thật tách rời.",
            "D": "Loại D (Nevertheless,) — Sai logic tương phản. Hai câu hoàn toàn bổ trợ và đồng thuận với nhau về thông tin lịch sử."
        }
    },
    25: {
        "id": 52,
        "master_id": "M52",
        "q_number": 25,
        "module_code": "RW_Module2",
        "topic": "Transitions & Clause-Level Logic (Contrast & Concession)",
        "difficulty": "Hard",
        "passage": "Between February 1876 and September 1876, George Eliot’s novel Daniel Deronda was serialized in eight parts in Blackwood’s Magazine. Scholars typically regard the serial version of Eliot’s novel as less authoritative than the version later published in book form. ______ as scholar Philip Allingham notes, serialized texts were the versions in which the nineteenth-century readership first “saw itself and its world reflected, and to which it responded, often with great emotion.”",
        "prompt": "Which choice completes the text with the most logical transition?",
        "options_raw": "A. Accordingly, they emphasize the appeal of the serial version for Victorian audiences;\nB. Still, the serial version is important to an understanding of the novel’s historical context;\nC. Consequently, the book version of the novel is often favored over the serial version;\nD. However, others regard the serial version as a more faithful representation of Eliot’s intentions;",
        "correct_answer": "B",
        "princeton_strategy": "Transitions - Concession & Historical Context",
        "tier1_hint": "Xác định sự đối lập: Câu trước nêu quan điểm của các học giả rằng bản đăng nhiều kỳ 'kém chuẩn xác/thẩm quyền hơn' (less authoritative). Câu sau trích dẫn Allingham chỉ ra rằng độc giả thế kỷ 19 lần đầu nhìn thấy thế giới của mình qua bản đăng báo này và phản hồi đầy cảm xúc.",
        "tier2_hint": "Cần một từ nối thể hiện sự nhượng bộ/tương phản (Still, However) kết hợp với nội dung chuẩn xác: khẳng định giá trị của bản đăng báo đối với việc hiểu bối cảnh lịch sử tiếp nhận tác phẩm.",
        "tier3_hint": "Phương án B 'Still, the serial version is important to an understanding of the novel’s historical context;' kết nối hoàn hảo: Dẫu vậy, bản đăng báo vẫn quan trọng để hiểu bối cảnh lịch sử. Chọn B.",
        "strategic_tip": "Mẹo Princeton Review: Cảnh giác với bẫy phương án D: Trích dẫn của Allingham nói về 'phản ứng của độc giả' (historical readership/context), KHÔNG hề nói về 'ý đồ tác giả' (Eliot's intentions).",
        "level1_tip": "Câu trước đánh giá bản đăng báo thấp hơn bản sách (less authoritative) -> Câu sau khẳng định giá trị lịch sử độc đáo của nó.",
        "level2_hint": "Cần từ nối nhượng bộ (Still) nói về bối cảnh lịch sử độc giả, không suy diễn sang ý đồ tác giả.",
        "level3_solution": "Chọn B. Still, the serial version is important to an understanding of the novel’s historical context;",
        "target_clue": "Đối lập giữa quan điểm học thuật thông thường ('less authoritative') và giá trị lịch sử thực tế trong việc phản ánh cảm xúc độc giả thế kỷ 19 ('saw itself and its world reflected').",
        "solution_steps": "1. Câu trước: Giới học giả coi bản đăng báo nhiều kỳ kém thẩm quyền hơn bản sách xuất bản sau đó.\n2. Câu sau: Học giả Allingham lưu ý độc giả thế kỷ 19 lần đầu nhìn thấy chính mình và phản hồi xúc động qua bản báo này -> Đây là giá trị bối cảnh lịch sử tiếp nhận.\n3. Cần chuyển tiếp nhượng bộ: Dù bị coi là kém thẩm quyền hơn, DẪU VẬY (Still) bản đăng báo vẫn quan trọng để hiểu bối cảnh lịch sử của tiểu thuyết.\n4. Loại A vì 'Accordingly' là hệ quả thuận, mâu thuẫn với việc câu trước đánh giá thấp. Loại C vì câu sau đang tôn vinh bản báo, không phải bản sách. Loại D là bẫy sai chi tiết: trích dẫn nói về độc giả, không nói về 'ý đồ của Eliot' (Eliot's intentions). Chọn B.",
        "pooe_breakdown": "• Đáp án ĐÚNG: B. Still, the serial version is important to an understanding of the novel’s historical context; — Sử dụng từ nối nhượng bộ 'Still' chính xác và liên kết trực tiếp với trích dẫn về độc giả lịch sử của Allingham.\n• Loại A. Accordingly... — Sai liên từ; 'Accordingly' biểu thị hệ quả thuận, không thể đi sau nhận xét 'less authoritative'.\n• Loại C. Consequently, the book version... — Sai trọng tâm; câu sau đang giải thích tầm quan trọng của bản đăng báo nhiều kỳ, không phải bản sách.\n• Loại D. However, others regard... Eliot’s intentions; — Bẫy sai đối tượng; đoạn văn nói về sự tiếp nhận của độc giả (readership), không nói về ý đồ sáng tác của Eliot.",
        "pro_tip": "Mẹo Princeton Review: Cảnh giác với bẫy phương án D: Trích dẫn của Allingham nói về 'phản ứng của độc giả' (historical readership/context), KHÔNG hề nói về 'ý đồ tác giả' (Eliot's intentions).",
        "option_explanations": {
            "A": "Loại A — Sai quan hệ logic. 'Accordingly' (Theo đó) là liên từ chỉ hệ quả thuận, không thể dùng khi câu trước đang đánh giá bản báo là 'less authoritative' (kém chuẩn mực hơn).",
            "B": "ĐÁP ÁN ĐÚNG: B — Chuẩn xác ngữ nghĩa và logic chuyển ý. 'Still' (Dẫu vậy) mở đầu sự nhượng bộ hợp lý: dù bị coi là kém chuẩn xác hơn bản sách, bản đăng báo vẫn đóng vai trò cốt lõi trong việc thấu hiểu bối cảnh lịch sử tiếp nhận tác phẩm.",
            "C": "Loại C — Lạc đề. Câu trích dẫn sau đó nói về giá trị cảm xúc của độc giả đối với bản đăng báo, không phải khẳng định việc bản sách được ưa chuộng hơn.",
            "D": "Loại D — Bẫy suy diễn sai lệch (Misleading Detail). Allingham chỉ nhấn mạnh trải nghiệm của độc giả thế kỷ 19 (bối cảnh lịch sử), hoàn toàn không đề cập đến việc đây có phải là 'ý đồ nghệ thuật của Eliot' hay không."
        }
    },
    26: {
        "id": 53,
        "master_id": "M53",
        "q_number": 26,
        "module_code": "RW_Module2",
        "topic": "Rhetorical Synthesis (Audience Adaptation / Introducing a Concept)",
        "difficulty": "Medium",
        "passage": "•  The Highpointers Club is a hiking club.\n•  One of the main goals among club members is to reach the highest points in all fifty US states.\n•  Those who achieve this are called 50 Completers.\n•  In Suk Han became a 50 Completer on November 15, 2013.\n•  The highest point in Wisconsin is Timms Hill, at 1,951 ft.\n•  The highest point in Wyoming is Gannett Peak, at 13,809 ft.",
        "prompt": "Which choice most effectively uses information from the given sentences to explain the 50 Completers hiking challenge to a new audience?",
        "options_raw": "A. If you are looking for a new hiking challenge, consider joining the Highpointers Club, as did In Suk Han, a hiker who successfully reached the highest point in every US state.\nB. Not until after you have reached the highest points in all fifty US states—including Timms Hill in Wisconsin and Gannett Peak in Wyoming—can you include yourself among the 50 Completers of the Highpointers Club.\nC. On November 15, 2013, In Suk Han finally completed the feat of reaching the highest point in all fifty US states, including Timms Hill in Wisconsin and Gannett Peak in Wyoming.\nD. Hikers aiming to count themselves among the 50 Completers must reach not only Wisconsin’s Timms Hill but also the even higher peak of Gannett Peak in Wyoming.",
        "correct_answer": "B",
        "princeton_strategy": "Rhetorical Synthesis - Goal Alignment & Concept Definition",
        "tier1_hint": "Đọc kỹ mục tiêu của đề bài: 'explain the 50 Completers hiking challenge to a new audience' (giải thích thử thách leo núi 50 Completers cho một độc giả mới).",
        "tier2_hint": "Câu trả lời đúng bắt buộc phải nêu rõ định nghĩa thử thách 50 Completers là gì (chinh phục đỉnh cao nhất của cả 50 bang của câu lạc bộ Highpointers Club), có thể kèm ví dụ minh họa.",
        "tier3_hint": "Phương án B định nghĩa trọn vẹn thử thách cho người mới: chinh phục điểm cao nhất của cả 50 bang nước Mỹ (ví dụ Timms Hill và Gannett Peak) để trở thành 50 Completer của Highpointers Club. Chọn B.",
        "strategic_tip": "Mẹo Princeton Review: Đối với câu hỏi Rhetorical Synthesis, hãy gạch chân mục tiêu đề bài (the student wants to...). Nếu mục tiêu là giải thích thử thách cho người mới, câu nào quá sa đà vào câu chuyện cá nhân một người (In Suk Han) sẽ bị loại!",
        "level1_tip": "Mục tiêu: Giải thích thử thách '50 Completers' cho độc giả mới.",
        "level2_hint": "Loại phương án tập trung vào cá nhân In Suk Han (A, C) hoặc thiếu thông tin 50 bang (D).",
        "level3_solution": "Chọn B. Not until after you have reached the highest points in all fifty US states...",
        "target_clue": "Mục tiêu đề bài: Giải thích thử thách leo núi '50 Completers' của câu lạc bộ Highpointers Club cho độc giả mới.",
        "solution_steps": "1. Xác định Goal: Giải thích thử thách '50 Completers' cho đối tượng người đọc mới.\n2. Yêu cầu nội dung: Phải nêu rõ leo đỉnh cao nhất của cả 50 bang nước Mỹ thuộc câu lạc bộ Highpointers Club.\n3. Phân tích các phương án:\n   - A: Kêu gọi tham gia câu lạc bộ và nói về In Suk Han, không giải thích rõ danh hiệu 50 Completers.\n   - C: Trọng tâm nói về sự kiện In Suk Han đạt danh hiệu vào năm 2013, không phải giải thích thử thách tổng quát cho người mới.\n   - D: Chỉ nhắc đến 2 đỉnh ở Wisconsin và Wyoming, bỏ sót điều kiện tiên quyết 'all fifty US states' và tên câu lạc bộ.\n   - B: Định nghĩa hoàn hảo bản chất thử thách (chinh phục điểm cao nhất của cả 50 bang, lấy ví dụ 2 đỉnh cụ thể) để được công nhận là 50 Completer của Highpointers Club. Chọn B.",
        "pooe_breakdown": "• Đáp án ĐÚNG: B. Not until after you have reached the highest points in all fifty US states... — Định nghĩa chính xác, đầy đủ bản chất thử thách 50 Completers cho người mới bắt đầu tìm hiểu.\n• Loại A. If you are looking... — Lạc trọng tâm sang việc rủ tham gia câu lạc bộ và nhắc tên cá nhân In Suk Han.\n• Loại C. On November 15, 2013... — Tập trung tường thuật chiến tích cá nhân của In Suk Han thay vì giải thích thể lệ thử thách.\n• Loại D. Hikers aiming... — Thiếu thông tin cốt lõi 'all fifty US states' và không nhắc đến câu lạc bộ chủ quản.",
        "pro_tip": "Mẹo Princeton Review: Đối với câu hỏi Rhetorical Synthesis, hãy gạch chân mục tiêu đề bài (the student wants to...). Nếu mục tiêu là giải thích thử thách cho người mới, câu nào quá sa đà vào câu chuyện cá nhân một người (In Suk Han) sẽ bị loại!",
        "option_explanations": {
            "A": "Loại A — Lệch mục tiêu. Câu này đóng vai trò lời mời gia nhập câu lạc bộ gắn với tên In Suk Han, không giải thích rõ ràng khái niệm thử thách 50 Completers là gì.",
            "B": "ĐÁP ÁN ĐÚNG: B — Đáp ứng trọn vẹn mục tiêu của đề bài. Giải thích rõ ràng điều kiện để trở thành 50 Completer của Highpointers Club (chinh phục đỉnh cao nhất của cả 50 bang Mỹ) kèm 2 ví dụ cụ thể.",
            "C": "Loại C — Sai trọng tâm. Câu này tập trung tường thuật mốc thời gian hoàn thành của cá nhân In Suk Han, không phục vụ mục đích giải thích thử thách cho độc giả mới.",
            "D": "Loại D — Thiếu dữ kiện then chốt. Chỉ nhắc đến hai đỉnh núi ở Wisconsin và Wyoming mà bỏ quên quy tắc cốt lõi là chinh phục toàn bộ 50 bang của Highpointers Club."
        }
    },
    27: {
        "id": 54,
        "master_id": "M54",
        "q_number": 27,
        "module_code": "RW_Module2",
        "topic": "Rhetorical Synthesis (Achieving a Specific Goal / Superiority Comparison)",
        "difficulty": "Hard",
        "passage": "While researching a topic, a student has taken the following notes:\n•  Enhanced rock weathering (ERW) is a method of spreading crushed basalt on farmland to capture carbon dioxide (CO2) and improve soil health.\n•  ERW increases crop yields and removes harmful CO2 from the air.\n•  ERW offers a low-cost, low-tech alternative to expensive, energy-intensive carbon capture methods.\n•  The CO2 captured by ERW becomes stable bicarbonate in soil water, which eventually streams into the ocean to remain safely stored for hundreds of years.\n•  ERW can improve rural economies by reducing farmers’ need for fertilizers and increasing their earnings from carbon credits.",
        "prompt": "The student wants to explain why ERW is a superior method for capturing CO2. Which choice most effectively uses relevant information from the notes to accomplish this goal?",
        "options_raw": "A. ERW improves soil health and helps farmers increase their income from selling carbon credits.\nB. ERW removes CO2 from the air in a way that is both low-cost and low-tech.\nC. The CO2 captured by ERW becomes stable bicarbonate, which improves soil health by remaining in the ocean for hundreds of years.\nD. ERW is a method of spreading crushed basalt on farmland to capture CO2, improving rural economies.",
        "correct_answer": "B",
        "princeton_strategy": "Rhetorical Synthesis - Goal Alignment & Comparative Superiority",
        "tier1_hint": "Đọc kỹ mục tiêu đề bài: 'explain why ERW is a superior method for capturing CO2' (giải thích vì sao ERW là phương pháp ƯU VIỆT HƠN để thu giữ CO2).",
        "tier2_hint": "Để chứng minh một phương pháp là 'superior' (vượt trội/ưu việt hơn), thông tin bắt buộc phải chứa sự so sánh với các phương pháp khác. Trong ghi chú, gạch đầu dòng nào so sánh ERW với các phương pháp khác?",
        "tier3_hint": "Ghi chú số 3 nêu rõ: 'ERW offers a low-cost, low-tech alternative to expensive, energy-intensive carbon capture methods'. Phương án B nêu trực tiếp ưu thế vượt trội này: thu giữ CO2 với chi phí thấp và công nghệ đơn giản ('low-cost and low-tech'). Chọn B.",
        "strategic_tip": "Mẹo Princeton Review: Khi đề bài yêu cầu chứng minh tính 'superior' (ưu việt hơn), câu trả lời bắt buộc phải phản ánh ưu thế cạnh tranh so với các giải pháp khác (ở đây là low-cost, low-tech thay vì tốn kém, ngốn điện).",
        "level1_tip": "Mục tiêu: Giải thích vì sao ERW ưu việt hơn trong việc thu giữ CO2.",
        "level2_hint": "Tìm thông tin so sánh: chi phí thấp và công nghệ đơn giản (low-cost, low-tech).",
        "level3_solution": "Chọn B. ERW removes CO2 from the air in a way that is both low-cost and low-tech.",
        "target_clue": "Mục tiêu: 'explain why ERW is a superior method for capturing CO2' gắn liền với ghi chú so sánh 'low-cost, low-tech alternative to expensive, energy-intensive methods'.",
        "solution_steps": "1. Xác định Goal: Giải thích lý do ERW là phương pháp thu giữ CO2 ưu việt hơn (superior).\n2. Đối chiếu các ghi chú:\n   - Ghi chú 3 là ghi chú duy nhất so sánh tính ưu việt của ERW với các phương pháp thu giữ carbon khác: nó mang lại giải pháp thay thế chi phí thấp (low-cost) và công nghệ đơn giản (low-tech) so với các phương pháp đắt đỏ và tốn năng lượng.\n3. Phân tích các lựa chọn:\n   - A: Nói về cải thiện đất và thu nhập nông dân, không giải thích vì sao thu CO2 ưu việt hơn.\n   - C: Mô tả cơ chế bicarbonate trong nước biển, không nêu tính ưu việt so sánh.\n   - D: Định nghĩa cách rải đá bazan và kinh tế nông thôn, không nói về tính ưu việt thu CO2.\n   - B: Nêu chuẩn xác: ERW thu gom CO2 bằng phương thức vừa có chi phí thấp vừa có công nghệ đơn giản (low-cost and low-tech). Chọn B.",
        "pooe_breakdown": "• Đáp án ĐÚNG: B. ERW removes CO2 from the air in a way that is both low-cost and low-tech. — Nêu trực tiếp ưu thế vượt trội của ERW trong việc thu giữ CO2 so với các công nghệ đắt đỏ khác.\n• Loại A. ERW improves soil health... — Lạc đề; tập trung vào lợi ích kinh tế và cải tạo đất của nông dân.\n• Loại C. The CO2 captured... — Mô tả tiến trình hóa học, không trả lời câu hỏi vì sao phương pháp này ưu việt hơn.\n• Loại D. ERW is a method... — Chỉ là câu giới thiệu định nghĩa tổng quát, không làm nổi bật tính ưu việt cạnh tranh.",
        "pro_tip": "Mẹo Princeton Review: Khi đề bài yêu cầu chứng minh tính 'superior' (ưu việt hơn), câu trả lời bắt buộc phải phản ánh ưu thế cạnh tranh so với các giải pháp khác (ở đây là low-cost, low-tech thay vì tốn kém, ngốn điện).",
        "option_explanations": {
            "A": "Loại A — Lệch mục tiêu. Câu này nói về sức khỏe của đất và thu nhập từ tín chỉ carbon của nông dân, không giải thích lý do vì sao ERW ưu việt hơn trong việc thu giữ CO2.",
            "B": "ĐÁP ÁN ĐÚNG: B — Khớp hoàn hảo với mục tiêu đề bài. Chỉ rõ tính ưu việt vượt trội của ERW khi thu giữ CO2: vừa có chi phí thấp (low-cost) vừa dùng công nghệ đơn giản (low-tech) thay vì các phương pháp đắt đỏ, tốn năng lượng.",
            "C": "Loại C — Không trả lời đúng trọng tâm. Câu này giải thích chu trình chuyển hóa bicarbonate vào đại dương, không nêu bật lý do phương pháp này vượt trội hơn.",
            "D": "Loại D — Thiếu thông tin so sánh. Đây chỉ là câu định nghĩa phương pháp rải đá bazan, không chỉ ra được tính ưu việt của ERW."
        }
    }
}

def update_json_file(file_path):
    with open(file_path, 'r', encoding='utf-8') as f:
        data = json.load(f)
    
    count = 0
    for q in data.get('questions', []):
        q_num = q.get('q_number')
        if q.get('module_code') == 'RW_Module2' and q_num in NEW_DATA_RW2:
            new_info = NEW_DATA_RW2[q_num]
            for k, v in new_info.items():
                q[k] = v
            count += 1
            print(f"[{file_path}] Updated Q{q_num} (ID {q.get('id')})")
            
    with open(file_path, 'w', encoding='utf-8') as f:
        json.dump(data, f, ensure_ascii=False, indent=2)
    print(f"Successfully saved {count} questions to {file_path}")

def update_excel_explanations(file_path):
    wb = openpyxl.load_workbook(file_path)
    
    # 1. Master Sheet
    if 'Master_All_98_Explanations' in wb.sheetnames:
        ws = wb['Master_All_98_Explanations']
        for r in range(5, 103):
            mid = ws.cell(r, 1).value
            mod = ws.cell(r, 2).value
            q_val = ws.cell(r, 3).value
            for q_num, info in NEW_DATA_RW2.items():
                if mid == info["master_id"] or (mod == 'RW_Module2' and q_val in [q_num, str(q_num), f"Q{q_num}"]):
                    ws.cell(r, 4).value = info["topic"]
                    ws.cell(r, 5).value = info["difficulty"]
                    ws.cell(r, 6).value = info["passage"]
                    ws.cell(r, 7).value = info["prompt"]
                    ws.cell(r, 8).value = info["options_raw"]
                    ws.cell(r, 9).value = info["correct_answer"]
                    ws.cell(r, 10).value = info["princeton_strategy"]
                    ws.cell(r, 11).value = info["target_clue"]
                    ws.cell(r, 12).value = info["solution_steps"]
                    ws.cell(r, 13).value = info["pooe_breakdown"]
                    ws.cell(r, 14).value = info["pro_tip"]
                    print(f"Updated Master_All_98_Explanations row {r} for {mid}")

    # 2. RW_Module2_Explanations Sheet
    if 'RW_Module2_Explanations' in wb.sheetnames:
        ws = wb['RW_Module2_Explanations']
        for r in range(5, ws.max_row + 1):
            q_val = ws.cell(r, 1).value
            for q_num, info in NEW_DATA_RW2.items():
                if q_val in [q_num, str(q_num), f"Q{q_num}"]:
                    ws.cell(r, 2).value = info["correct_answer"]
                    ws.cell(r, 3).value = info["target_clue"]
                    ws.cell(r, 4).value = info["solution_steps"]
                    ws.cell(r, 5).value = info["pooe_breakdown"]
                    ws.cell(r, 6).value = info["topic"]
                    ws.cell(r, 7).value = info["difficulty"]
                    ws.cell(r, 8).value = info["passage"]
                    ws.cell(r, 9).value = info["prompt"]
                    ws.cell(r, 10).value = info["options_raw"]
                    ws.cell(r, 11).value = info["princeton_strategy"]
                    ws.cell(r, 12).value = info["pro_tip"]
                    print(f"Updated RW_Module2_Explanations row {r} for Q{q_num}")

    wb.save(file_path)
    print(f"Successfully saved {file_path}")

def update_excel_template(file_path):
    wb = openpyxl.load_workbook(file_path)
    if 'RW_Module2' in wb.sheetnames:
        ws = wb['RW_Module2']
        for r in range(1, ws.max_row + 1):
            q_val = ws.cell(r, 1).value
            for q_num, info in NEW_DATA_RW2.items():
                if q_val in [q_num, str(q_num), f"Q{q_num}"]:
                    ws.cell(r, 2).value = info["topic"]
                    ws.cell(r, 3).value = info["passage"]
                    ws.cell(r, 4).value = info["prompt"]
                    ws.cell(r, 5).value = info["options_raw"]
                    print(f"Updated {file_path} sheet RW_Module2 row {r} for Q{q_num}")
                    
    if 'Master_All_98_Questions' in wb.sheetnames:
        ws = wb['Master_All_98_Questions']
        for r in range(1, ws.max_row + 1):
            mid = ws.cell(r, 1).value
            for q_num, info in NEW_DATA_RW2.items():
                if mid == info["master_id"]:
                    ws.cell(r, 3).value = info["topic"]
                    ws.cell(r, 4).value = info["passage"]
                    ws.cell(r, 5).value = info["prompt"]
                    ws.cell(r, 6).value = info["options_raw"]
                    print(f"Updated {file_path} sheet Master_All_98_Questions row {r} for {mid}")
                    
    wb.save(file_path)
    print(f"Successfully saved {file_path}")

if __name__ == '__main__':
    update_json_file('SAT Coach/BluebookOct2026App/database.json')
    update_json_file('SAT Coach/BluebookOct2026App/static/database.json')
    update_excel_explanations('SAT Coach/Real Tests/Oct 2026/Explanations.xlsx')
    update_excel_template('SAT Coach/BluebookOct2026App/template.xlsx')
    update_excel_template('SAT Coach/BluebookOct2026App/static/template.xlsx')
