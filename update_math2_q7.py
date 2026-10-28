import json
import openpyxl

NEW_DATA = {
    "passage": "A laboratory culture of bacteria grows according to the model P(t) = 250(1.08)^t, where t represents the elapsed time in hours.",
    "prompt": "Calculator Allowed. By what percentage does the population increase each hour?",
    "options_raw": "A. 0.08%\nB. 8%\nC. 25%\nD. 108%",
    "correct_answer": "B",
    "topic": "Exponential Population Growth Rate",
    "tier1_hint": "Mô hình tăng trưởng hàm số mũ có dạng P(t) = P0 * (1 + r)^t, trong đó r là tỉ lệ tăng trưởng.",
    "tier2_hint": "So sánh hệ số trong ngoặc (1.08) với biểu thức (1 + r): 1 + r = 1.08 => r = 0.08.",
    "tier3_hint": "Đổi số thập phân r = 0.08 sang tỉ lệ phần trăm bằng cách nhân với 100%: 0.08 × 100% = 8%.",
    "strategic_tip": "Growth Factor = 1 + r => 1.08 = 1 + 0.08 => r = 8%.",
    "level1_tip": "Xác định tỉ lệ tăng trưởng trong mô hình tăng trưởng mũ P(t) = P0(1 + r)^t.",
    "level2_hint": "So sánh hệ số cơ số 1.08 với dạng (1 + r).",
    "level3_solution": "1 + r = 1.08 => r = 0.08 = 8%.",
    "target_clue": "Mô hình hàm số mũ P(t) = 250(1.08)^t có dạng P(t) = P0 * (1 + r)^t.\nHệ số cơ số (growth factor) là 1.08 = 1 + 0.08, nghĩa là tỉ lệ tăng trưởng r = 0.08 = 8% mỗi giờ.",
    "solution_steps": "Mô hình hàm số mũ chuẩn: P(t) = P0 * (1 + r)^t.\nSo sánh với P(t) = 250(1.08)^t:\n1) Số lượng vi khuẩn ban đầu P0 = 250.\n2) Hệ số cơ số (growth factor) 1 + r = 1.08 => r = 1.08 - 1 = 0.08.\n3) Tỉ lệ phần trăm tăng trưởng mỗi giờ: r = 0.08 × 100% = 8%.",
    "princeton_strategy": "Exponential Growth Rate",
    "pro_tip": "Mẹo Princeton Review: Đối với mô hình hàm số mũ P(t) = P0 * b^t, phần trăm tăng trưởng mỗi đơn vị thời gian luôn bằng (b - 1) × 100%. Ở đây: (1.08 - 1) × 100% = 8%.",
    "pooe_breakdown": "• Đáp án ĐÚNG: B. 8% — 1 + r = 1.08 => r = 0.08 = 8% mỗi giờ.\n• Loại A. 0.08% — Lỗi chưa nhân với 100% (0.08 = 8%, không phải 0.08%).\n• Loại C. 25% — Lỗi nhầm lẫn từ hệ số ban đầu 250.\n• Loại D. 108% — 108% là quy mô vi khuẩn sau 1 giờ so với ban đầu (1 + r), không phải phần trăm tăng thêm r.",
    "option_explanations": {
        "A": "Loại A (0.08%) — Lỗi nhầm lẫn do chưa nhân với 100% khi chuyển số thập phân 0.08 sang tỉ lệ phần trăm (0.08 = 8%, không phải 0.08%).",
        "B": "ĐÁP ÁN ĐÚNG: B (8%) — Mô hình tăng trưởng hàm số mũ chuẩn có dạng P(t) = P0 * (1 + r)^t. Với hệ số cơ số trong ngoặc là 1.08 = 1 + 0.08, tỉ lệ tăng trưởng vi khuẩn mỗi giờ là r = 0.08 = 8%.",
        "C": "Loại C (25%) — Lỗi nhầm lẫn lấy từ hệ số ban đầu 250 (hoặc 25% từ 250).",
        "D": "Loại D (108%) — Lỗi lấy toàn bộ hệ số cơ số nhân 100% (108% là tỉ lệ dân số mới so với ban đầu, không phải phần trăm TĂNG THÊM r)."
    }
}

def update_json_file(file_path):
    with open(file_path, 'r', encoding='utf-8') as f:
        data = json.load(f)
    
    updated = False
    for q in data.get('questions', []):
        if q.get('id') == 83 or (q.get('module_code') == 'Math_Module2' and q.get('q_number') == 7):
            for k, v in NEW_DATA.items():
                q[k] = v
            updated = True
            print(f"Updated {file_path} for question ID {q.get('id')}")
            break
            
    if updated:
        with open(file_path, 'w', encoding='utf-8') as f:
            json.dump(data, f, ensure_ascii=False, indent=2)
        print(f"Successfully saved {file_path}")
    else:
        print(f"Warning: Question not found in {file_path}")

def update_excel(file_path):
    wb = openpyxl.load_workbook(file_path)
    
    # 1. Master Sheet
    if 'Master_All_98_Explanations' in wb.sheetnames:
        ws = wb['Master_All_98_Explanations']
        for r in range(2, ws.max_row + 1):
            if ws.cell(r, 1).value == 'M83' or (ws.cell(r, 2).value == 'Math_Module2' and ws.cell(r, 3).value in [7, '7', 'Q7']):
                # Col 4: Topic
                ws.cell(r, 4).value = NEW_DATA["topic"]
                # Col 6: Passage
                ws.cell(r, 6).value = NEW_DATA["passage"]
                # Col 7: Prompt
                ws.cell(r, 7).value = NEW_DATA["prompt"]
                # Col 8: Options
                ws.cell(r, 8).value = NEW_DATA["options_raw"]
                # Col 9: Correct Answer
                ws.cell(r, 9).value = NEW_DATA["correct_answer"]
                # Col 10: Strategy
                ws.cell(r, 10).value = NEW_DATA["princeton_strategy"]
                # Col 11: Target & Clue
                ws.cell(r, 11).value = NEW_DATA["target_clue"]
                # Col 12: Solution Steps
                ws.cell(r, 12).value = NEW_DATA["solution_steps"]
                # Col 13: POOE Breakdown
                ws.cell(r, 13).value = NEW_DATA["pooe_breakdown"]
                # Col 14: Pro-Tip
                ws.cell(r, 14).value = NEW_DATA["pro_tip"]
                print(f"Updated Master_All_98_Explanations row {r}")
                break

    # 2. Math_Module2_Explanations
    if 'Math_Module2_Explanations' in wb.sheetnames:
        ws = wb['Math_Module2_Explanations']
        for r in range(2, ws.max_row + 1):
            if ws.cell(r, 1).value in ['Q7', 7, '7']:
                # Col 2: Key
                ws.cell(r, 2).value = NEW_DATA["correct_answer"]
                # Col 3: Pro Tip
                ws.cell(r, 3).value = NEW_DATA["pro_tip"]
                # Col 4: Topic
                ws.cell(r, 4).value = NEW_DATA["topic"]
                # Col 6: Passage
                ws.cell(r, 6).value = NEW_DATA["passage"]
                # Col 7: Prompt
                ws.cell(r, 7).value = NEW_DATA["prompt"]
                # Col 8: Options
                ws.cell(r, 8).value = NEW_DATA["options_raw"]
                # Col 9: Strategy
                ws.cell(r, 9).value = NEW_DATA["princeton_strategy"]
                # Col 10: Target & Clue
                ws.cell(r, 10).value = NEW_DATA["target_clue"]
                # Col 11: Solution Steps
                ws.cell(r, 11).value = NEW_DATA["solution_steps"]
                # Col 12: POOE Breakdown
                ws.cell(r, 12).value = NEW_DATA["pooe_breakdown"]
                print(f"Updated Math_Module2_Explanations row {r}")
                break

    wb.save(file_path)
    print(f"Successfully saved Excel: {file_path}")

if __name__ == '__main__':
    update_json_file('SAT Coach/BluebookOct2026App/database.json')
    update_json_file('SAT Coach/BluebookOct2026App/static/database.json')
    update_excel('SAT Coach/Real Tests/Oct 2026/Explanations.xlsx')
