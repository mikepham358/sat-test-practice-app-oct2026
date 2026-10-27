import os
import re
import json
import fitz
import openpyxl

APP_DIR = "/Users/mike/Library/CloudStorage/GoogleDrive-mikepham358@gmail.com/My Drive/A Vinschool by Mike Pham/SAT Coach/BluebookOct2026App"
KEY_PDF_PATH = "/Users/mike/Library/CloudStorage/GoogleDrive-mikepham358@gmail.com/My Drive/A Vinschool by Mike Pham/SAT Coach/Real Tests/Oct 2026/KeyB.pdf"

DB_PATH = os.path.join(APP_DIR, "database.json")
STATIC_DB_PATH = os.path.join(APP_DIR, "static", "database.json")
TEMPLATE_PATH = os.path.join(APP_DIR, "template.xlsx")
STATIC_TEMPLATE_PATH = os.path.join(APP_DIR, "static", "template.xlsx")
MASTER_TEMPLATE_PATH = "/Users/mike/Library/CloudStorage/GoogleDrive-mikepham358@gmail.com/My Drive/A Vinschool by Mike Pham/26-27 Grand Plan /Study Plan/Module /Template - October 2026.xlsx"

SEC_MAPPING = {
    "Section 1, Module 1: Reading and Writing": "RW_Module1",
    "Section 1, Module 2: Reading and Writing": "RW_Module2",
    "Section 2, Module 1: Math": "Math_Module1",
    "Section 2, Module 2: Math": "Math_Module2"
}

def extract_keys_from_pdf():
    doc = fitz.open(KEY_PDF_PATH)
    full_text = doc[0].get_text() + "\n" + doc[1].get_text()
    sections = re.split(r"(Section \d, Module \d: [^\n]+)", full_text)
    
    official_keys = {}
    for i in range(1, len(sections), 2):
        sec_header = sections[i].strip()
        sec_content = sections[i+1].strip()
        mod_code = SEC_MAPPING[sec_header]
        
        lines = [l.strip() for l in sec_content.splitlines() if l.strip()]
        clean_tokens = [
            l for l in lines 
            if l not in ["Question #", "Correct", "Answer", "Answer Key"]
            and "sattest.hocsat.com" not in l 
            and "Digital SAT Practice Test" not in l 
            and "For student-produced response" not in l 
            and "Detailed explanations" not in l
        ]
        
        pairs = {}
        idx = 0
        while idx < len(clean_tokens):
            tok = clean_tokens[idx]
            if re.match(r"^\d+$", tok) and idx + 1 < len(clean_tokens):
                pairs[int(tok)] = clean_tokens[idx+1].strip()
                idx += 2
            else:
                idx += 1
        official_keys[mod_code] = pairs
        
    total_q = sum(len(p) for p in official_keys.values())
    print(f"Extracted {total_q} official keys from KeyB.pdf across {len(official_keys)} modules.")
    assert total_q == 98, f"Expected 98 keys, got {total_q}"
    return official_keys

def update_database(official_keys):
    print("Updating database.json and static/database.json...")
    with open(DB_PATH, "r", encoding="utf-8") as f:
        db = json.load(f)

    updated_keys = 0
    for q in db["questions"]:
        mod = q["module_code"]
        qnum = q["q_number"]
        if mod in official_keys and qnum in official_keys[mod]:
            off_ans = official_keys[mod][qnum]
            old_ans = q.get("correct_answer")
            if old_ans != off_ans:
                print(f"  Fixing [{mod}] Q{qnum}: {old_ans} -> {off_ans}")
                updated_keys += 1
            
            q["correct_answer"] = off_ans
            is_mcq = (off_ans in ["A", "B", "C", "D"])
            q["is_grid_in"] = not is_mcq
            
            if is_mcq:
                q["options"] = [
                    {"key": "A", "text": "Lựa chọn A"},
                    {"key": "B", "text": "Lựa chọn B"},
                    {"key": "C", "text": "Lựa chọn C"},
                    {"key": "D", "text": "Lựa chọn D"}
                ]
                q["options_raw"] = "A. Lựa chọn A\nB. Lựa chọn B\nC. Lựa chọn C\nD. Lựa chọn D"
            else:
                q["options"] = []
                q["options_raw"] = "Student-Produced Response (Grid-in)"

    with open(DB_PATH, "w", encoding="utf-8") as f:
        json.dump(db, f, ensure_ascii=False, indent=2)

    with open(STATIC_DB_PATH, "w", encoding="utf-8") as f:
        json.dump(db, f, ensure_ascii=False, indent=2)

    print(f"database.json saved. Changed {updated_keys} answer keys.")

def update_excel_workbooks(official_keys):
    workbooks = [TEMPLATE_PATH, STATIC_TEMPLATE_PATH, MASTER_TEMPLATE_PATH]
    
    for path in workbooks:
        if not os.path.exists(path):
            print(f"Warning: {path} not found, skipping.")
            continue
            
        print(f"Updating formulas & options in {path}...")
        wb = openpyxl.load_workbook(path)
        
        # 1. Update 4 Module Sheets
        for mod, pairs in official_keys.items():
            if mod in wb.sheetnames:
                ws = wb[mod]
                for qnum, off_ans in pairs.items():
                    r = 4 + qnum
                    is_mcq = (off_ans in ["A", "B", "C", "D"])
                    
                    # Col 5: Options
                    if is_mcq:
                        ws.cell(r, 5).value = "A. Lựa chọn A\nB. Lựa chọn B\nC. Lựa chọn C\nD. Lựa chọn D"
                    else:
                        ws.cell(r, 5).value = "Student-Produced Response (Grid-in)"
                        
                    # Col 7: Formula L1
                    ws.cell(r, 7).value = f'=IF(ISBLANK(F{r}), "⏳ Chờ đáp án", IF(UPPER(TRIM(F{r}))="{off_ans}", "✅ Đúng", "❌ Sai"))'
                    # Col 9: Formula L2
                    ws.cell(r, 9).value = f'=IF(ISBLANK(H{r}), "", IF(UPPER(TRIM(H{r}))="{off_ans}", "✅ Đúng", "❌ Sai"))'
                    # Col 11: Formula L3
                    ws.cell(r, 11).value = f'=IF(ISBLANK(J{r}), "", IF(UPPER(TRIM(J{r}))="{off_ans}", "✅ Đúng", "❌ Sai"))'

        # 2. Update Master_All_98_Questions Sheet
        if "Master_All_98_Questions" in wb.sheetnames:
            ws_m = wb["Master_All_98_Questions"]
            
            # Row mapping:
            # RW1: rows 5 to 31 (M1 to M27)
            # RW2: rows 32 to 58 (M28 to M54)
            # Math1: rows 59 to 80 (M55 to M76)
            # Math2: rows 81 to 102 (M77 to M98)
            mod_ranges = [
                ("RW_Module1", 1, 27, 5),
                ("RW_Module2", 1, 27, 32),
                ("Math_Module1", 1, 22, 59),
                ("Math_Module2", 1, 22, 81),
            ]
            for mod, start_q, end_q, start_row in mod_ranges:
                pairs = official_keys[mod]
                for qnum in range(start_q, end_q + 1):
                    r = start_row + (qnum - start_q)
                    off_ans = pairs[qnum]
                    is_mcq = (off_ans in ["A", "B", "C", "D"])
                    
                    # Col 6: Options
                    if is_mcq:
                        ws_m.cell(r, 6).value = "A. Lựa chọn A\nB. Lựa chọn B\nC. Lựa chọn C\nD. Lựa chọn D"
                    else:
                        ws_m.cell(r, 6).value = "Student-Produced Response (Grid-in)"
                        
                    # Col 8: Formula L1
                    ws_m.cell(r, 8).value = f'=IF(ISBLANK(G{r}), "⏳ Chờ đáp án", IF(UPPER(TRIM(G{r}))="{off_ans}", "✅ Đúng", "❌ Sai"))'
                    # Col 10: Formula L2
                    ws_m.cell(r, 10).value = f'=IF(ISBLANK(I{r}), "", IF(UPPER(TRIM(I{r}))="{off_ans}", "✅ Đúng", "❌ Sai"))'
                    # Col 12: Formula L3
                    ws_m.cell(r, 12).value = f'=IF(ISBLANK(K{r}), "", IF(UPPER(TRIM(K{r}))="{off_ans}", "✅ Đúng", "❌ Sai"))'

        wb.save(path)
        print(f"Successfully saved {path}.")

def verify_results(official_keys):
    print("\n--- Verifying Results ---")
    with open(DB_PATH, "r", encoding="utf-8") as f:
        db = json.load(f)

    db_mismatches = 0
    for q in db["questions"]:
        mod = q["module_code"]
        qnum = q["q_number"]
        db_ans = str(q.get("correct_answer")).strip().upper()
        off_ans = str(official_keys[mod][qnum]).strip().upper()
        if db_ans != off_ans:
            print(f"Error in DB: {mod} Q{qnum} {db_ans} != {off_ans}")
            db_mismatches += 1

    print(f"Verification Database Mismatches: {db_mismatches} (Must be 0)")
    assert db_mismatches == 0, "Database still has mismatches!"

if __name__ == "__main__":
    official_keys = extract_keys_from_pdf()
    update_database(official_keys)
    update_excel_workbooks(official_keys)
    verify_results(official_keys)
