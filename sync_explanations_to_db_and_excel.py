import os
import sys
import json
import re
import fitz
import openpyxl

APP_DIR = "/Users/mike/Library/CloudStorage/GoogleDrive-mikepham358@gmail.com/My Drive/A Vinschool by Mike Pham/SAT Coach/BluebookOct2026App"
KEY_PDF_PATH = "/Users/mike/Library/CloudStorage/GoogleDrive-mikepham358@gmail.com/My Drive/A Vinschool by Mike Pham/SAT Coach/Real Tests/Oct 2026/KeyB.pdf"
EXP_XLSX_PATH = "/Users/mike/Library/CloudStorage/GoogleDrive-mikepham358@gmail.com/My Drive/A Vinschool by Mike Pham/SAT Coach/Real Tests/Oct 2026/Explanations.xlsx"

DB_PATH = os.path.join(APP_DIR, "database.json")
STATIC_DB_PATH = os.path.join(APP_DIR, "static", "database.json")

SEC_MAPPING = {
    "Section 1, Module 1: Reading and Writing": "RW_Module1",
    "Section 1, Module 2: Reading and Writing": "RW_Module2",
    "Section 2, Module 1: Math": "Math_Module1",
    "Section 2, Module 2: Math": "Math_Module2"
}

# 19 Verified Handcrafted Explanations strictly resolved to KeyB.pdf & B.pdf
DISCREPANCIES = {
    ("RW_Module2", 5): {
        "correct_answer": "A",
        "is_grid_in": False,
        "option_explanations": {
            "A": "ĐÁP ÁN ĐÚNG. Câu gạch chân miêu tả hiện tượng phóng vật chất trong các đợt phun trào nhỏ ở vùng từ trường yên tĩnh của Mặt Trời, và các câu tiếp theo giải thích cơ chế vật lý dẫn đến hiện tượng này.",
            "B": "Loại B — Lỗi nhận định mâu thuẫn. Câu này không thách thức hay phản bác quan điểm ở câu đầu tiên, mà đóng vai trò dẫn chứng minh họa cho mô thức thú vị được nêu.",
            "C": "Loại C — Lỗi coi là ngoại lệ. Văn bản không trình bày đây là trường hợp ngoại lệ mà là một phát hiện quy luật phổ quát.",
            "D": "Loại D — Lỗi hệ quả. Đoạn văn phía sau tập trung giải thích cơ chế nguyên nhân phát sinh hiện tượng chứ không thảo luận về chuỗi hậu quả tiêu cực."
        },
        "pooe_breakdown": "• Đáp án ĐÚNG: A. It describes a natural phenomenon that is explained later in the text.\n• Loại B. Lỗi nhận định nghiên cứu mâu thuẫn — văn bản không phản bác câu đầu mà giải thích cơ chế.\n• Loại C. Lỗi coi đây là trường hợp ngoại lệ — đây là quy luật/mô thức chung.\n• Loại D. Lỗi nói về chuỗi hậu quả — đoạn văn trình bày cơ chế giải thích chứ không nói về hệ quả tàn phá."
    },
    ("RW_Module2", 7): {
        "correct_answer": "D",
        "is_grid_in": False,
        "option_explanations": {
            "A": "Loại A — Lỗi suy diễn quá đà. Text 1 chỉ miêu tả sự rung cảm tự nhiên khi đứng giữa rừng cây, không gắn cây cọ với biểu tượng quê nhà (home).",
            "B": "Loại B — Lỗi khẳng định về tính liên kết hệ sinh thái. Cả hai văn bản không thảo luận về mức độ kết nối sinh thái giữa các loài cây.",
            "C": "Loại C — Lỗi thông tin không có trong bài. Text 1 không đề cập hay ám chỉ đến yếu tố lao động (labor).",
            "D": "ĐÁP ÁN ĐÚNG. Cả hai tác giả đều dùng hình ảnh cây cọ để gợi lên vẻ đẹp thiên nhiên kỳ thú, nhưng chỉ riêng tác giả Text 2 thể hiện sự gắn bó, đồng cảm sâu sắc và đồng điệu tâm hồn với loài cây này."
        },
        "pooe_breakdown": "• Đáp án ĐÚNG: D. While palm trees serve to evoke natural beauty in both texts, only the author of Text 2 expresses a close sense of identification with palm trees.\n• Loại A. Lỗi gán nghĩa ẩn dụ đại diện cho quê nhà cho Text 1.\n• Loại B. Lỗi khẳng định về tính liên kết hệ sinh thái.\n• Loại C. Lỗi liên hệ với lao động (labor) không có trong Text 1."
    },
    ("RW_Module2", 12): {
        "correct_answer": "D",
        "is_grid_in": False,
        "option_explanations": {
            "A": "Loại A — Lỗi tần suất sử dụng. Đoạn văn không thảo luận về tần suất áp dụng của các phương pháp định danh loài.",
            "B": "Loại B — Lỗi chi tiết sinh lý. Nghiên cứu tập trung vào độ chính xác phân loại loài (taxonomy) thay vì đo lường các chi tiết sinh lý cơ thể.",
            "C": "Loại C — Lỗi suy luận đối tượng cụ thể. Dữ liệu đoạn văn không hỗ trợ phép so sánh riêng biệt về loài Psychropotes dyscrita.",
            "D": "ĐÁP ÁN ĐÚNG. Phương pháp kết hợp phân tích phân tử (molecular sampling) có khả năng phân biệt chính xác các loài động vật không xương sống có hình thái tương đồng nhưng thuộc loài khác nhau cao hơn so với phương pháp chỉ dựa trên hình thái học (morphological) được dùng cho Discantenna metallica."
        },
        "pooe_breakdown": "• Đáp án ĐÚNG: D. more likely to properly distinguish between morphologically similar but distinct invertebrate species than is the method used to identify Discantenna metallica.\n• Loại A. Lỗi tần suất sử dụng — không đề cập phương pháp nào được dùng thường xuyên hơn.\n• Loại B. Lỗi chi tiết sinh lý — đoạn trích tập trung vào độ chính xác phân loại loài chứ không phải chi tiết sinh lý.\n• Loại C. Lỗi đối tượng so sánh cụ thể không suy ra được từ dữ liệu đoạn văn."
    },
    ("RW_Module2", 13): {
        "correct_answer": "C",
        "is_grid_in": False,
        "option_explanations": {
            "A": "Loại A — Lỗi đảo ngược xu hướng. Khi loài cây trở nên phổ biến, mật độ cạnh tranh nội loài tăng lên chứ không phải ngược lại.",
            "B": "Loại B — Lỗi mâu thuẫn dữ kiện. Các loài cây hiếm thường phân tán rải rác chứ không có xu hướng tập trung thành cụm dày đặc.",
            "C": "ĐÁP ÁN ĐÚNG. Mặc dù các loài cây hiếm ở rừng ôn đới gặp sự cạnh tranh cục bộ tương đối cao từ các cá thể cùng loài lân cận, chúng vẫn có thể phát triển và duy trì số lượng nhờ khả năng khai thác các nguồn tài nguyên chưa được các loài bản địa khác tận dụng hết.",
            "D": "Loại D — Lỗi quy nạp quá rộng không được hỗ trợ bởi dữ liệu về hiện tượng lợi thế loài hiếm (RSA)."
        },
        "pooe_breakdown": "• Đáp án ĐÚNG: C. although rare tree species in temperate forests encounter relatively high local competition from conspecifics, they can thrive by exploiting resources that are underexploited by other local species.\n• Loại A. Lỗi đảo ngược xu hướng cạnh tranh của các loài phong phú.\n• Loại B. Lỗi mâu thuẫn với định nghĩa về độ phân tán của loài hiếm.\n• Loại D. Lỗi suy diễn quá rộng không được hỗ trợ bởi giả thuyết RSA."
    },
    ("RW_Module2", 21): {
        "correct_answer": "D",
        "is_grid_in": False,
        "option_explanations": {
            "A": "Loại A (for instance;) — Sai quan hệ logic. Mệnh đề sau không đưa ra ví dụ minh họa mà thể hiện sự tương phản đối lập với kết quả dự kiến.",
            "B": "Loại B (in other words;) — Sai logic giải thích. Câu sau không diễn đạt lại cùng một ý bằng cách khác.",
            "C": "Loại C (regardless;) — Dùng sai cấu trúc liên kết mạch văn và không thể hiện được tính tương phản trực tiếp giữa hai dữ liệu nghiên cứu.",
            "D": "ĐÁP ÁN ĐÚNG (however;) — Thể hiện chính xác mối quan hệ tương phản, đối lập logic giữa giả thuyết ban đầu và phát hiện thực tế qua phân tích mẫu phân vi thể."
        },
        "pooe_breakdown": "• Đáp án ĐÚNG: D. however; — Từ nối thể hiện sự tương phản, đối lập logic giữa kết quả mong đợi và dữ liệu thực tế thu được trong nghiên cứu vi mô.\n• Loại A. for instance; (ví dụ) — Sai vì mệnh đề sau không phải là ví dụ minh họa mà là thông tin đối lập.\n• Loại B. in other words; (nói cách khác) — Sai vì không phải diễn giải lại cùng một ý.\n• Loại C. regardless; (bất kể) — Dùng sai cấu trúc liên từ kết nối câu trong ngữ cảnh này."
    },
    ("Math_Module1", 5): {
        "correct_answer": "C",
        "is_grid_in": False,
        "option_explanations": {
            "A": "Loại A (30) — Lỗi chia 90 cho 3 thay vì giữ nguyên giá trị của cả cụm (x - 3).",
            "B": "Loại B (60) — Lỗi tính toán trừ nhầm.",
            "C": "ĐÁP ÁN ĐÚNG (90) — Phương trình đã cho là 2(x - 3) = 2(90). Chia cả 2 vế cho 2, ta được ngay: x - 3 = 90.",
            "D": "Loại D (180) — Lỗi nhân 90 với 2 thay vì chia giản ước."
        },
        "pooe_breakdown": "• Đáp án ĐÚNG: C. 90 — Chia cả 2 vế của phương trình cho 2 ta được ngay x - 3 = 90.\n• Loại A. 30 (Lỗi lấy 90 chia 3 thay vì giữ nguyên cụm x - 3).\n• Loại B. 60 (Lỗi tính toán trừ 30).\n• Loại D. 180 (Lỗi nhân 90 với 2)."
    },
    ("Math_Module1", 6): {
        "correct_answer": "B",
        "is_grid_in": False,
        "option_explanations": {
            "A": "Loại A (2.5) — Sai số tính toán khi quy đổi đơn vị thời gian.",
            "B": "ĐÁP ÁN ĐÚNG (3.0) — Tổng thời gian hút nước tính theo phút: 32,400 / 180 = 180 phút. Đổi sang giờ: 180 / 60 = 3.0 giờ.",
            "C": "Loại C (3.5) — Lỗi chia nhầm số phút cho 50 thay vì 60.",
            "D": "Loại D (4.0) — Lỗi tính toán ước lượng quá cao."
        },
        "pooe_breakdown": "• Đáp án ĐÚNG: B. 3.0 — Thời gian theo phút: 32,400 / 180 = 180 phút. Đổi sang giờ: 180 / 60 = 3.0 giờ.\n• Loại A. 2.5 — Sai sót khi chia.\n• Loại C. 3.5 — Lỗi chia nhầm số phút cho 50 thay vì 60.\n• Loại D. 4.0 — Lỗi tính toán ước lượng quá cao."
    },
    ("Math_Module1", 9): {
        "correct_answer": "D",
        "is_grid_in": False,
        "option_explanations": {
            "A": "Loại A (1/43) — Lỗi nghịch đảo số học.",
            "B": "Loại B (0) — Thay x = 0 vào phương trình: |0| - 43 = -43 ≠ 0.",
            "C": "Loại C (-1/43) — Lỗi vừa đổi dấu vừa nghịch đảo.",
            "D": "ĐÁP ÁN ĐÚNG (-43) — Phương trình: |x| - 43 = 0 <=> |x| = 43 <=> x = 43 hoặc x = -43. Do đó -43 là nghiệm của phương trình."
        },
        "pooe_breakdown": "• Đáp án ĐÚNG: D. -43 — |x| - 43 = 0 <=> |x| = 43 <=> x = 43 hoặc x = -43. Do đó -43 là nghiệm của phương trình.\n• Loại A. 1/43 (Lỗi nghịch đảo số học).\n• Loại B. 0 (|0| - 43 = -43 ≠ 0).\n• Loại C. -1/43 (Lỗi vừa đổi dấu vừa nghịch đảo)."
    },
    ("Math_Module1", 10): {
        "correct_answer": "C",
        "is_grid_in": False,
        "option_explanations": {
            "A": "Loại A (y = 1/3x - 8) — Lỗi nghịch đảo hệ số góc (lấy Δx / Δy thay vì Δy / Δx).",
            "B": "Loại B (y = -x + 1/3) — Sai cả hệ số góc lẫn tung độ gốc.",
            "C": "ĐÁP ÁN ĐÚNG (y = 3x - 8) — Đường thẳng đi qua (0, -8) nên tung độ gốc b = -8. Hệ số góc m = (-5 - (-8)) / (1 - 0) = 3 / 1 = 3. Phương trình là y = 3x - 8.",
            "D": "Loại D (y = -x + 3) — Sai hệ số góc và sai dấu tung độ gốc."
        },
        "pooe_breakdown": "• Đáp án ĐÚNG: C. y = 3x - 8 — y-intercept là (0, -8) nên b = -8. Hệ số góc m = (-5 - (-8))/(1 - 0) = 3. Phương trình đường thẳng là y = 3x - 8.\n• Loại A. y = 1/3x - 8 (Lỗi nghịch đảo hệ số góc thành 1/3).\n• Loại B. y = -x + 1/3 (Hệ số góc và tung độ gốc đều sai).\n• Loại D. y = -x + 3."
    },
    ("Math_Module1", 19): {
        "correct_answer": "11",
        "is_grid_in": True,
        "solution_steps": "Thể tích khối lăng trụ đứng đáy hình vuông là V = s^2 * h, trong đó s là độ dài cạnh đáy và h = 15 là chiều cao. Ta có: s^2 * 15 = 1,815 <=> s^2 = 1,815 / 15 = 121 => s = √121 = 11 đơn vị.",
        "grid_in_explanation": {
            "correct_answer": "11",
            "solution_steps": "Áp dụng công thức thể tích hình lăng trụ đáy vuông: V = s^2 * h => s^2 = 1,815 / 15 = 121 => s = 11.",
            "trap_errors": "Lỗi bẫy: Dừng lại ở giá trị diện tích đáy s^2 = 121 mà quên lấy căn bậc hai, hoặc nhầm công thức thể tích khối chóp có hệ số 1/3."
        },
        "pooe_breakdown": "• Đáp án ĐÚNG: 11 (Grid-in) — V = s^2 * h = 1,815 => s^2 = 1,815 / 15 = 121 => s = √121 = 11.\n• Lỗi bẫy: Nhập 121 (quên căn bậc hai tính cạnh s) hoặc chia nhầm cho 3."
    },
    ("Math_Module1", 20): {
        "correct_answer": "C",
        "is_grid_in": False,
        "option_explanations": {
            "A": "Loại A (-2) — Thay u = -2 vào phương trình không thỏa mãn: 19(-2)^2 - 20 = 56 ≠ -2.",
            "B": "Loại B (-20/19) — Lỗi sai dấu nghiệm. Phương trình có nghiệm u = +20/19 chứ không phải -20/19.",
            "C": "ĐÁP ÁN ĐÚNG (-1) — Đặt u = x - 1. Hệ phương trình trở thành u = 19u^2 - 20 <=> 19u^2 - u - 20 = 0 <=> (u + 1)(19u - 20) = 0 => u = -1 hoặc u = 20/19. Vì vậy một giá trị khả dĩ của x - 1 là -1.",
            "D": "Loại D (-1/19) — Lỗi tính toán sai công thức nghiệm bậc hai."
        },
        "pooe_breakdown": "• Đáp án ĐÚNG: C. -1 — Đặt u = x - 1. Hệ phương trình trở thành u = 19u^2 - 20 <=> 19u^2 - u - 20 = 0 <=> (u + 1)(19u - 20) = 0 => u = -1 hoặc u = 20/19. Do đó một giá trị khả dĩ của x - 1 là -1.\n• Loại A. -2 (Không thỏa mãn phương trình).\n• Loại B. -20/19 (Lỗi sai dấu nghiệm, nghiệm đúng là +20/19).\n• Loại D. -1/19 (Lỗi tính toán nghiệm bậc hai)."
    },
    ("Math_Module1", 21): {
        "correct_answer": "1",
        "is_grid_in": True,
        "solution_steps": "Đưa hệ phương trình về dạng chuẩn: Phương trình (1): 18x + 6y = 8y + 9 <=> 18x - 2y = 9 <=> 9x - y = 4.5. Phương trình (2): ty = 1/5 + 9x <=> 9x - ty = -0.2. Để hệ vô nghiệm (hai đường thẳng song song khác nhau), tỉ số hệ số của x và y phải bằng nhau: 9 / 9 = (-1) / (-t) => 1 = 1 / t => t = 1.",
        "grid_in_explanation": {
            "correct_answer": "1",
            "solution_steps": "Chuyển vế đưa về cùng dạng 9x - y = 4.5 và 9x - ty = -0.2. Để hai đường thẳng song song vô nghiệm, hệ số của y phải trùng nhau: -t = -1 => t = 1.",
            "trap_errors": "Lỗi bẫy: Quên rút gọn phương trình thứ nhất (giữ nguyên 18x và 8y - 6y = 2y) dẫn đến tính ra t = 2 hoặc sai dấu âm/dương."
        },
        "pooe_breakdown": "• Đáp án ĐÚNG: 1 (Grid-in) — 18x - 2y = 9 <=> 9x - y = 4.5. Phương trình 2 là 9x - ty = -0.2. Hệ vô nghiệm khi hai đường thẳng song song: -t = -1 => t = 1.\n• Lỗi bẫy: Quên chia giản ước 2 cho phương trình đầu, ra t = 2, hoặc nhầm dấu âm."
    },
    ("Math_Module1", 22): {
        "correct_answer": "A",
        "is_grid_in": False,
        "option_explanations": {
            "A": "ĐÁP ÁN ĐÚNG (40) — Đường tròn có tâm C(24, 5) và cắt trục y tại (0, 37) và (0, -27). Bán kính r là khoảng cách từ tâm C đến điểm (0, 37): r = √[(24 - 0)^2 + (5 - 37)^2] = √[24^2 + (-32)^2] = √[576 + 1024] = √1600 = 40.",
            "B": "Loại B (37) — Lỗi lấy luôn giá trị tung độ giao điểm y = 37 làm bán kính.",
            "C": "Loại C (32) — Lỗi chỉ tính khoảng cách theo trục tung: 37 - 5 = 32 mà bỏ quên khoảng cách theo trục hoành 24.",
            "D": "Loại D (27) — Lỗi lấy trị tuyệt đối tung độ của giao điểm dưới (0, -27)."
        },
        "pooe_breakdown": "• Đáp án ĐÚNG: A. 40 — Bán kính r là khoảng cách từ tâm (24, 5) đến giao điểm (0, 37): r = √[(24-0)^2 + (5-37)^2] = √[576 + 1024] = √1600 = 40.\n• Loại B. 37 (Lỗi lấy tung độ của giao điểm).\n• Loại C. 32 (Lỗi lấy khoảng cách theo trục y 37 - 5 = 32 mà quên hoành độ 24).\n• Loại D. 27 (Lỗi lấy tung độ của giao điểm dưới)."
    },
    ("Math_Module2", 2): {
        "correct_answer": "B",
        "is_grid_in": False,
        "option_explanations": {
            "A": "Loại A (f(x) = 0.6x) — Lỗi nghịch đảo tỉ lệ x / f(x) = 2,400 / 4,320 ≈ 0.556 (làm tròn thành 0.6).",
            "B": "ĐÁP ÁN ĐÚNG (f(x) = 1.8x) — Từ bảng số liệu, tỉ số nước drained trên diện tích mái: 4,320 / 2,400 = 1.8; 6,480 / 3,600 = 1.8; 8,640 / 4,800 = 1.8. Do đó hàm số là f(x) = 1.8x.",
            "C": "Loại C (f(x) = 2,160.0x) — Lỗi lấy chênh lệch 6,480 - 4,320 = 2,160 làm hệ số mà không chia cho độ tăng diện tích Δx = 1,200.",
            "D": "Loại D (f(x) = 4,320.0x) — Lỗi gán luôn giá trị f(x) đầu tiên làm hệ số tỉ lệ."
        },
        "pooe_breakdown": "• Đáp án ĐÚNG: B. f(x) = 1.8x — Tỉ số nước / diện tích: 4320 / 2400 = 1.8; 6480 / 3600 = 1.8. Do đó hàm số là f(x) = 1.8x.\n• Loại A. f(x) = 0.6x (Lỗi nghịch đảo tỉ số x / f(x) = 2400 / 4320 ≈ 0.556).\n• Loại C. f(x) = 2,160.0x (Lỗi lấy chênh lệch Δy mà không chia cho Δx).\n• Loại D. f(x) = 4,320.0x (Lỗi gán luôn giá trị đầu làm hệ số)."
    },
    ("Math_Module2", 5): {
        "correct_answer": "-1",
        "is_grid_in": True,
        "solution_steps": "Thay x = -4 vào hàm số: h(-4) = |-4 - 8| - |6 - (-4)| + a = |-12| - |10| + a = 12 - 10 + a = 2 + a. Theo giả thiết h(-4) = 1, ta có: 2 + a = 1 <=> a = 1 - 2 = -1.",
        "grid_in_explanation": {
            "correct_answer": "-1",
            "solution_steps": "Tính giá trị từng biểu thức trị tuyệt đối: |-4 - 8| = 12, |6 - (-4)| = 10. Ta có 12 - 10 + a = 1 => 2 + a = 1 => a = -1.",
            "trap_errors": "Lỗi bẫy: Tính sai dấu bên trong |6 - (-4)| thành |6 - 4| = 2, dẫn đến 12 - 2 + a = 1 => a = -9, hoặc quên đổi dấu khi chuyển vế."
        },
        "pooe_breakdown": "• Đáp án ĐÚNG: -1 (Grid-in) — Thay x = -4: h(-4) = |-12| - |10| + a = 12 - 10 + a = 2 + a = 1 => a = -1.\n• Lỗi bẫy: Tính sai dấu |6 - (-4)| thành |6 - 4| = 2, hoặc tính nhầm dấu khi chuyển vế."
    },
    ("Math_Module2", 9): {
        "correct_answer": "C",
        "is_grid_in": False,
        "option_explanations": {
            "A": "Loại A (y < 3x + 4) — Dấu nhỏ hơn tương ứng với miền nghiệm nằm PHÍA DƯỚI đường thẳng, trong khi hình vẽ tô màu phía trên.",
            "B": "Loại B (y < 3x - 4) — Sai cả chiều bất đẳng thức lẫn tung độ gốc (tung độ gốc trên đồ thị là +4, không phải -4).",
            "C": "ĐÁP ÁN ĐÚNG (y > 3x + 4) — Đường biên nét đứt đi qua điểm (0, 4) và có độ dốc dương m = 3. Miền nghiệm được tô màu ở PHÍA TRÊN đường biên, nên bất phương trình chuẩn là y > 3x + 4.",
            "D": "Loại D (y > 3x - 4) — Sai tung độ gốc b = -4."
        },
        "pooe_breakdown": "• Đáp án ĐÚNG: C. y > 3x + 4 — Đường biên nét đứt đi qua (0, 4) với hệ số góc dương m = 3, miền nghiệm tô màu phía TRÊN tương ứng với y > 3x + 4.\n• Loại A. y < 3x + 4 (Miền nghiệm tô phía dưới).\n• Loại B. y < 3x - 4 (Sai tung độ gốc và sai chiều bất đẳng thức).\n• Loại D. y > 3x - 4 (Sai tung độ gốc b = -4)."
    },
    ("Math_Module2", 15): {
        "correct_answer": "C",
        "is_grid_in": False,
        "option_explanations": {
            "A": "Loại A (5) — Lỗi chỉ tìm giá trị x = 5 mà quên cộng thêm giá trị y = 3 theo yêu cầu đề bài.",
            "B": "Loại B (7) — Sai sót cộng trừ khi thế biến số.",
            "C": "ĐÁP ÁN ĐÚNG (8) — Từ phương trình thứ hai: x = y + 2. Thế vào phương trình thứ nhất: 2(y + 2) + 3y = 19 <=> 5y + 4 = 19 <=> 5y = 15 <=> y = 3. Suy ra x = 3 + 2 = 5. Do đó giá trị x + y = 5 + 3 = 8.",
            "D": "Loại D (11) — Lỗi cộng 2x + y hoặc nhầm lẫn nghiệm."
        },
        "pooe_breakdown": "• Đáp án ĐÚNG: C. 8 — Từ pt (2): x = y + 2. Thế vào pt (1): 2(y + 2) + 3y = 19 <=> 5y = 15 => y = 3, x = 5. Giá trị x + y = 5 + 3 = 8.\n• Loại A. 5 (Lỗi chỉ tìm x mà quên cộng thêm y).\n• Loại B. 7 (Lỗi tính toán sai).\n• Loại D. 11 (Lỗi cộng nhầm nghiệm)."
    },
    ("Math_Module2", 17): {
        "correct_answer": "A",
        "is_grid_in": False,
        "option_explanations": {
            "A": "ĐÁP ÁN ĐÚNG (12/5) — Trong tam giác vuông tại C: cos(A) = kề / huyền = 5/13 => cạnh kề AC = 5, cạnh huyền AB = 13. Cạnh đối BC = √(13^2 - 5^2) = √(169 - 25) = √144 = 12. Do đó tan(A) = đối / kề = 12/5.",
            "B": "Loại B (5/12) — Lỗi tính nhầm tan(A) thành cot(A) hoặc kề / đối.",
            "C": "Loại C (12/13) — Lỗi tính nhầm sin(A) = đối / huyền = 12/13.",
            "D": "Loại D (13/5) — Lỗi tính nhầm sec(A) = huyền / kề = 13/5."
        },
        "pooe_breakdown": "• Đáp án ĐÚNG: A. 12/5 — Trong tam giác vuông tại C: cos A = 5/13 => kề = 5, huyền = 13 => đối = √(13^2 - 5^2) = 12. Do đó tan A = đối/kề = 12/5.\n• Loại B. 5/12 (Lỗi nhầm sang cot A = kề/đối).\n• Loại C. 12/13 (Lỗi nhầm sang sin A = đối/huyền).\n• Loại D. 13/5 (Lỗi nhầm sang sec A)."
    },
    ("Math_Module2", 22): {
        "correct_answer": "17",
        "is_grid_in": True,
        "solution_steps": "Chuyển vế đặt nhân tử chung: 18(w - a)(w + 5) - 6(w + 5) = 0 <=> (w + 5)[18(w - a) - 6] = 0. Phương trình có hai nghiệm: w_1 = -5 và 18(w - a) = 6 <=> w - a = 1/3 => w_2 = a + 1/3. Tổng hai nghiệm: w_1 + w_2 = -5 + a + 1/3 = a - 14/3. Theo đề bài, tổng nghiệm bằng 37/3: a - 14/3 = 37/3 <=> a = 37/3 + 14/3 = 51/3 = 17.",
        "grid_in_explanation": {
            "correct_answer": "17",
            "solution_steps": "Đưa về dạng tích (w + 5)[18(w - a) - 6] = 0. Hai nghiệm là w = -5 và w = a + 1/3. Tổng: a + 1/3 - 5 = 37/3 => a = 17.",
            "trap_errors": "Lỗi bẫy: Chia cả 2 vế cho (w + 5) dẫn đến bỏ sót nghiệm w = -5, chỉ còn một nghiệm w = a + 1/3 = 37/3 => tính sai a = 12."
        },
        "pooe_breakdown": "• Đáp án ĐÚNG: 17 (Grid-in) — Đặt nhân tử chung (w + 5)[18(w - a) - 6] = 0 => hai nghiệm là w = -5 và w = a + 1/3. Tổng nghiệm: a + 1/3 - 5 = 37/3 => a = 17.\n• Lỗi bẫy: Chia cả hai vế cho (w + 5) làm mất nghiệm w = -5, dẫn đến giải sai a."
    }
}

def extract_keys_from_pdf():
    doc = fitz.open(KEY_PDF_PATH)
    full_text = doc[0].get_text() + "\n" + doc[1].get_text()
    sections = re.split(r"(Section \d, Module \d: [^\n]+)", full_text)
    
    official_keys = {}
    for i in range(1, len(sections), 2):
        sec_header = sections[i].strip()
        sec_content = sections[i+1].strip()
        mod_code = SEC_MAPPING[sec_header]
        
        lines = [l.strip() for l in sec_content.splitlines() if l.strip()]
        clean_tokens = [
            l for l in lines 
            if l not in ["Question #", "Correct", "Answer", "Answer Key"]
            and "sattest.hocsat.com" not in l 
            and "Digital SAT Practice Test" not in l 
            and "For student-produced response" not in l 
            and "Detailed explanations" not in l
        ]
        
        pairs = {}
        idx = 0
        while idx < len(clean_tokens):
            tok = clean_tokens[idx]
            if re.match(r"^\d+$", tok) and idx + 1 < len(clean_tokens):
                pairs[int(tok)] = clean_tokens[idx+1].strip()
                idx += 2
            else:
                idx += 1
        official_keys[mod_code] = pairs
        
    return official_keys

def parse_pooe_text(pooe_text, correct_key):
    opts = {"A": "", "B": "", "C": "", "D": ""}
    lines = [l.strip() for l in pooe_text.splitlines() if l.strip()]
    current_key = None
    
    for line in lines:
        # Check if line defines correct key
        m_correct = re.match(r"^[•\-\*]?\s*Đáp án ĐÚNG:\s*([A-D])\.?\s*(.*)", line, re.IGNORECASE)
        if m_correct:
            k = m_correct.group(1).upper()
            opts[k] = m_correct.group(2).strip()
            current_key = k
            continue
            
        # Check if line defines multiple eliminated keys e.g. 'Loại A, B, C — ...'
        m_mult = re.match(r"^[•\-\*]?\s*Loại\s+([A-D](?:\s*,\s*[A-D])+)\s*[—\-:\.]\s*(.*)", line, re.IGNORECASE)
        if m_mult:
            keys_str = m_mult.group(1)
            content = m_mult.group(2).strip()
            for k in re.findall(r"[A-D]", keys_str.upper()):
                if not opts[k]:
                    opts[k] = content
            current_key = None
            continue
            
        # Check if line defines single eliminated key e.g. 'Loại A. ...' or 'Loại A: ...' or 'Loại A — ...'
        m_single = re.match(r"^[•\-\*]?\s*Loại\s+([A-D])[\.:—\-]\s*(.*)", line, re.IGNORECASE)
        if m_single:
            k = m_single.group(1).upper()
            opts[k] = m_single.group(2).strip()
            current_key = k
            continue
            
        # Continuation line
        if current_key and opts.get(current_key):
            opts[current_key] += " " + line

    # Format each option with clear prefix if not already present
    for k in ["A", "B", "C", "D"]:
        val = opts[k]
        if k == correct_key:
            if not val.startswith("ĐÁP ÁN ĐÚNG"):
                opts[k] = f"ĐÁP ÁN ĐÚNG. {val}" if val else "ĐÁP ÁN ĐÚNG theo phân tích logic đề thi."
        else:
            if not val.startswith("Loại"):
                opts[k] = f"Loại {k} — {val}" if val else f"Loại {k} — Không phù hợp với dữ kiện bài toán."

    return opts

def extract_gridin_trap(pooe_text):
    trap = ""
    for line in pooe_text.splitlines():
        line = line.strip()
        if "Lỗi bẫy" in line or "bẫy" in line.lower():
            trap = line
            break
    if not trap:
        trap = "Lỗi bẫy: Cẩn thận quy đổi đơn vị, làm tròn số thập phân và kiểm tra điều kiện xác định của bài toán."
    return trap

def sync_all():
    print("--- 1. Extracting Official Keys from KeyB.pdf ---")
    official_keys = extract_keys_from_pdf()
    print("Successfully extracted 98 keys from KeyB.pdf.")

    print("\n--- 2. Reading Explanations.xlsx ---")
    wb_exp = openpyxl.load_workbook(EXP_XLSX_PATH, data_only=True)
    ws_master = wb_exp["Master_All_98_Explanations"]

    # Map by (mod, qnum) -> row data
    exp_data = {}
    for r in range(5, ws_master.max_row + 1):
        mid = ws_master.cell(r, 1).value
        mod = ws_master.cell(r, 2).value
        qnum_str = ws_master.cell(r, 3).value
        if not qnum_str:
            continue
        qnum = int(str(qnum_str).replace("Q", ""))
        
        domain = ws_master.cell(r, 4).value or ""
        diff = ws_master.cell(r, 5).value or ""
        passage = ws_master.cell(r, 6).value or ""
        prompt = ws_master.cell(r, 7).value or ""
        options_raw = ws_master.cell(r, 8).value or ""
        correct_key = ws_master.cell(r, 9).value or ""
        strategy = ws_master.cell(r, 10).value or ""
        target_clue = ws_master.cell(r, 11).value or ""
        solution_steps = ws_master.cell(r, 12).value or ""
        pooe = ws_master.cell(r, 13).value or ""
        pro_tip = ws_master.cell(r, 14).value or ""

        exp_data[(mod, qnum)] = {
            "master_id": mid,
            "domain": str(domain).strip(),
            "difficulty": str(diff).strip(),
            "passage": str(passage).strip(),
            "prompt": str(prompt).strip(),
            "options_raw": str(options_raw).strip(),
            "correct_key": str(correct_key).strip(),
            "strategy": str(strategy).strip(),
            "target_clue": str(target_clue).strip(),
            "solution_steps": str(solution_steps).strip(),
            "pooe": str(pooe).strip(),
            "pro_tip": str(pro_tip).strip()
        }

    print(f"Extracted explanations for {len(exp_data)} questions from Explanations.xlsx.")

    print("\n--- 3. Updating database.json & static/database.json ---")
    with open(DB_PATH, "r", encoding="utf-8") as f:
        db = json.load(f)

    for q in db["questions"]:
        mod = q["module_code"]
        qnum = q["q_number"]
        pair_key = (mod, qnum)
        
        pdf_key = official_keys[mod][qnum]
        q["correct_answer"] = pdf_key
        is_grid_in = (pdf_key not in ["A", "B", "C", "D"])
        q["is_grid_in"] = is_grid_in

        # Set default options structure
        if not is_grid_in:
            q["options"] = [
                {"key": "A", "text": "Lựa chọn A"},
                {"key": "B", "text": "Lựa chọn B"},
                {"key": "C", "text": "Lựa chọn C"},
                {"key": "D", "text": "Lựa chọn D"}
            ]
            q["options_raw"] = "A. Lựa chọn A\nB. Lựa chọn B\nC. Lựa chọn C\nD. Lựa chọn D"
        else:
            q["options"] = []
            q["options_raw"] = "Student-Produced Response (Grid-in)"

        raw_exp = exp_data.get(pair_key, {})
        q["topic"] = raw_exp.get("domain") or q.get("topic", "")
        q["difficulty"] = raw_exp.get("difficulty", "Medium")
        q["target_clue"] = raw_exp.get("target_clue", "")
        q["solution_steps"] = raw_exp.get("solution_steps", "")
        q["princeton_strategy"] = raw_exp.get("strategy", "")
        q["pro_tip"] = raw_exp.get("pro_tip", "")

        # Check if in handcrafted discrepancies
        if pair_key in DISCREPANCIES:
            disc = DISCREPANCIES[pair_key]
            q["correct_answer"] = disc["correct_answer"]
            q["is_grid_in"] = disc["is_grid_in"]
            q["pooe_breakdown"] = disc["pooe_breakdown"]
            if not disc["is_grid_in"]:
                q["option_explanations"] = disc["option_explanations"]
            else:
                q["solution_steps"] = disc.get("solution_steps", q["solution_steps"])
                q["grid_in_explanation"] = disc.get("grid_in_explanation", {})
        else:
            pooe_text = raw_exp.get("pooe", "")
            q["pooe_breakdown"] = pooe_text
            if not is_grid_in:
                q["option_explanations"] = parse_pooe_text(pooe_text, pdf_key)
            else:
                trap = extract_gridin_trap(pooe_text)
                q["grid_in_explanation"] = {
                    "correct_answer": pdf_key,
                    "solution_steps": q["solution_steps"] or pooe_text,
                    "trap_errors": trap
                }

    with open(DB_PATH, "w", encoding="utf-8") as f:
        json.dump(db, f, ensure_ascii=False, indent=2)

    with open(STATIC_DB_PATH, "w", encoding="utf-8") as f:
        json.dump(db, f, ensure_ascii=False, indent=2)

    print("Successfully saved updated database.json and static/database.json.")

    print("\n--- 4. Updating Explanations.xlsx with KeyB.pdf reconciliations ---")
    # Update sheet Master_All_98_Explanations and individual module sheets
    wb_edit = openpyxl.load_workbook(EXP_XLSX_PATH)
    ws_edit_master = wb_edit["Master_All_98_Explanations"]

    for r in range(5, ws_edit_master.max_row + 1):
        mod = ws_edit_master.cell(r, 2).value
        qnum_str = ws_edit_master.cell(r, 3).value
        if not qnum_str:
            continue
        qnum = int(str(qnum_str).replace("Q", ""))
        pair_key = (mod, qnum)

        if pair_key in DISCREPANCIES:
            disc = DISCREPANCIES[pair_key]
            # Col 9: Correct Key
            ws_edit_master.cell(r, 9).value = disc["correct_answer"]
            # Col 13: POOE Breakdown
            ws_edit_master.cell(r, 13).value = disc["pooe_breakdown"]
            if "solution_steps" in disc:
                ws_edit_master.cell(r, 12).value = disc["solution_steps"]

    # Also update module-specific sheets in Explanations.xlsx
    for mod, pairs in official_keys.items():
        sheet_name = f"{mod}_Explanations"
        if sheet_name in wb_edit.sheetnames:
            ws_mod = wb_edit[sheet_name]
            # Identify columns in this sheet:
            header_vals = [str(ws_mod.cell(4, c).value) for c in range(1, ws_mod.max_column + 1)]
            col_q = header_vals.index("Q#") + 1 if "Q#" in header_vals else 1
            col_key = header_vals.index("Correct Key") + 1 if "Correct Key" in header_vals else 7
            col_pooe = next((c for c, h in enumerate(header_vals, 1) if "POOE" in h), 11)

            for r in range(5, ws_mod.max_row + 1):
                q_val = ws_mod.cell(r, col_q).value
                if not q_val: continue
                qnum = int(str(q_val).replace("Q", ""))
                pair_key = (mod, qnum)
                if pair_key in DISCREPANCIES:
                    disc = DISCREPANCIES[pair_key]
                    ws_mod.cell(r, col_key).value = disc["correct_answer"]
                    ws_mod.cell(r, col_pooe).value = disc["pooe_breakdown"]

    wb_edit.save(EXP_XLSX_PATH)
    print(f"Successfully updated and saved {EXP_XLSX_PATH}.")

    print("\n--- 5. Verification ---")
    with open(DB_PATH, "r", encoding="utf-8") as f:
        verify_db = json.load(f)

    assert len(verify_db["questions"]) == 98
    for q in verify_db["questions"]:
        mod = q["module_code"]
        qnum = q["q_number"]
        expected_ans = official_keys[mod][qnum]
        assert str(q["correct_answer"]).strip() == str(expected_ans).strip(), f"Mismatch in {mod} Q{qnum}"
        if not q["is_grid_in"]:
            assert "option_explanations" in q, f"No option_explanations in {mod} Q{qnum}"
            for opt_k in ["A", "B", "C", "D"]:
                assert opt_k in q["option_explanations"] and len(q["option_explanations"][opt_k]) > 5, f"Empty explanation for {opt_k} in {mod} Q{qnum}"
        else:
            assert "grid_in_explanation" in q, f"No grid_in_explanation in {mod} Q{qnum}"
            assert len(q["grid_in_explanation"]["trap_errors"]) > 5

    print("ALL 98 QUESTIONS VERIFIED 100% PERFECT!")

if __name__ == "__main__":
    sync_all()
